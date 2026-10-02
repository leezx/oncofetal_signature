#!/usr/bin/env python3
"""Export docs/DATASET_REGISTRY.md tables to CSV and a formatted XLSX.

Usage: python3 scripts/export_dataset_registry.py   (from the repo root)
Outputs: docs/DATASET_REGISTRY.csv, docs/DATASET_REGISTRY.xlsx
"""
import pathlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "DATASET_REGISTRY.md"


def md_tables(text):
    tables, cur = [], []
    for line in text.splitlines():
        if line.startswith("|"):
            cur.append(line)
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    out = []
    for t in tables:
        rows = [[c.strip().replace("**", "").replace("`", "") for c in r.strip().strip("|").split("|")]
                for r in t if not set(r.replace("|", "").strip()) <= set("-: ")]
        out.append(pd.DataFrame(rows[1:], columns=rows[0]))
    return out


def split_pmid_doi(df):
    pd_col = df.pop("PMID / DOI").str.split(" / ", n=1, expand=True)
    pos = df.columns.get_loc("Journal, year") + 1
    df.insert(pos, "PMID", pd_col[0].replace("—", ""))
    df.insert(pos + 1, "DOI", pd_col[1].fillna(""))
    return df


def write_sheet(ws, df, widths):
    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", start_color="1F4E78")
    body_font = Font(name="Arial", size=10)
    step_fill = {"2": "DDEBF7", "2 (benchmark)": "F2F2F2", "3": "E2EFDA"}
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for c in ws[1]:
        c.font, c.fill, c.border = head_font, head_fill, border
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in ws.iter_rows(min_row=2):
        fill = step_fill.get(str(r[0].value)) if "Step" in df.columns else None
        for c in r:
            c.font, c.border = body_font, border
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                c.fill = PatternFill("solid", start_color=fill)
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 18)
    ws.freeze_panes = "C2" if "Step" in df.columns else "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    main_df, unused = md_tables(MD.read_text())
    main_df = split_pmid_doi(main_df)
    main_df.to_csv(ROOT / "docs" / "DATASET_REGISTRY.csv", index=False)

    widths = {"Step": 12, "Role": 22, "Dataset": 26, "Source publication": 24, "Journal, year": 20,
              "PMID": 11, "DOI": 26, "Species": 11, "Accession": 30, "Assay / platform": 32,
              "Material & contrast": 36, "Units used": 22, "Status / notes": 34}
    wb = Workbook()
    ws = wb.active
    ws.title = "Datasets"
    write_sheet(ws, main_df, widths)
    write_sheet(wb.create_sheet("Not used or deferred"), unused,
                {"Dataset": 34, "Intended role": 40, "Reason": 60})
    notes = wb.create_sheet("Notes")
    for line in [
        "Source: docs/DATASET_REGISTRY.md (compiled 2026-10-02 from dataset link.md files, step reports, GEO series records and PubMed).",
        "Units used = biological replicates entering each contrast after the frozen eligibility rules.",
        "Row colour: blue = Step 2, grey = Step 2 benchmark, green = Step 3.",
        "Joanito (Synapse syn26844071) terms: non-commercial use; no transfer or disclosure of the data or derived material.",
        "GSE230581 and GSE160449 are sub-series of the same Pikkupeura 2023 study (not independent).",
        "Regenerate: python3 scripts/export_dataset_registry.py",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    for r in notes.iter_rows():
        r[0].font = Font(name="Arial", size=10)
    wb.save(ROOT / "docs" / "DATASET_REGISTRY.xlsx")
    print(f"{len(main_df)} datasets, {len(unused)} deferred")


if __name__ == "__main__":
    main()
