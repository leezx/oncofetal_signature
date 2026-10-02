#!/usr/bin/env python3
"""Summary statistics of the 31 Step 1 positive oncofetal markers across every
Step 2 (fetal vs adult) and Step 3 (cancer vs normal) dataset.

One row per marker; per dataset: log2FC, P value and FDR where the dataset's
frozen analysis produced them. Positive log2FC = higher in fetal (Step 2) or
in tumour/malignant (Step 3). Descriptive only: no gate is computed here.

Usage (repo root): python3 marker_summary/scripts/build_marker_summary.py

Writes the full (restricted, git-ignored) tables and a `_public` version that
omits Joanito columns and Joanito-derived Step 3 labels, which may be committed.
"""
import pathlib

import h5py
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA")
S2 = ROOT / "step2_fetal/results/tables"
BM = ROOT / "step2_dataset_benchmark/results/tables"
S3 = ROOT / "step3_cancer/results/tables"
OUT = ROOT / "marker_summary/results"

# Mouse genes used when no Ensembl 116 one-to-one orthologue exists.
MOUSE_OVERRIDE = {
    "LY6A": ("Ly6a", "mouse-defined marker (human LY6S one2many, conf 0)"),
    "REG3B": ("Reg3b", "mouse-defined marker (human REG1B one2many, conf 0)"),
    "CXADR": ("Cxadr", "Ensembl one2many; high-confidence Cxadr used"),
    "SPRR1A": (None, "no mouse orthologue in Ensembl 116"),
}

# (dataset key, step, label, species, contrast, positive direction)
DATASETS = [
    ("HGCA", "2", "HGCA (Elmentaite 2021)", "Human", "fetal vs adult epithelium (H1)", "fetal"),
    ("HGCA_prolif", "2", "HGCA control", "Human", "fetal vs adult stem/TA/progenitor", "fetal"),
    ("Gao", "2", "Gao 2018 (GSE103239)", "Human", "fetal vs adult large-intestinal epithelium (H2, effect only)", "fetal"),
    ("Mouse_GSE230581", "2", "Pikkupeura in vivo (GSE230581)", "Mouse", "E16.5 vs adult crypt epithelium (M1)", "fetal"),
    ("Senger", "2 (benchmark)", "Senger 2018 (GSE101531)", "Human", "fetal vs adult enterospheres", "fetal"),
    ("Fawkner", "2 (benchmark)", "Fawkner-Corbett 2021 (GSE158702)", "Human", "fetal EpCAM+ only: median detection fraction (no FC)", "n/a"),
    ("Pikkupeura_LN", "2 (benchmark)", "Pikkupeura cultures LN (GSE160449)", "Mouse", "fetal vs adult culture, laminin arm", "fetal"),
    ("Pikkupeura_collagen", "2 (benchmark)", "Pikkupeura cultures collagen (GSE160449)", "Mouse", "fetal vs adult culture, collagen arm", "fetal"),
    ("GSE44433", "2 (benchmark)", "Hemmerling 2014 (GSE44433)", "Mouse", "WT E17.5 vs WT 8-week LCM ileal epithelium", "fetal"),
    ("TCGA", "3", "TCGA COAD+READ (T1)", "Human", "primary tumour vs normal, ~ tissue + project", "tumour"),
    ("TCGA_paired", "3", "TCGA paired", "Human", "tumour vs matched normal (50 pairs)", "tumour"),
    ("TCGA_GTEx", "3", "TCGA vs GTEx", "Human", "TCGA tumour vs GTEx transverse colon (direction only)", "tumour"),
    ("Joanito", "3", "Joanito 2022 (S1)", "Human", "malignant vs normal epithelium, ~ cohort + group", "malignant"),
    ("Joanito_sens", "3", "Joanito sensitivity", "Human", "S1 restricted to cohorts with both groups", "malignant"),
    ("Pelka", "3", "Pelka 2021 (S2, GSE178341)", "Human", "tumour vs normal epithelium", "tumour"),
    ("Pelka_P", "3", "Pelka progenitor check (P)", "Human", "tumour epithelium vs normal stem/TA (cE01-03)", "tumour"),
]
STATS = ["log2FC", "PValue", "FDR"]


def lookup(df, key_col, keys, cols, status_universe=None):
    """Return {gene: (values..., status)} trying each key in order."""
    idx = df.drop_duplicates(key_col).set_index(key_col)
    out = {}
    for gene, cands in keys.items():
        hit = next((k for k in cands if k in idx.index), None)
        if hit is not None and pd.notna(idx.loc[hit, cols[0]]):
            out[gene] = [idx.loc[hit, c] if c else np.nan for c in cols] + ["measured"]
        elif status_universe is not None and any(k in status_universe for k in cands):
            out[gene] = [np.nan] * len(cols) + ["filtered (low expression)"]
        else:
            out[gene] = [np.nan] * len(cols) + ["not in dataset annotation"]
    return out


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cand = pd.read_csv(ROOT / "step2_fetal/config/literature_candidates_31.tsv", sep="\t")
    alias = pd.read_csv(ROOT / "step3_cancer/config/symbol_aliases.tsv", sep="\t")
    to_src = dict(zip(alias.current_hgnc_symbol, alias.source_symbol))
    human_keys = {g: [g] + ([to_src[g]] if g in to_src else []) for g in cand.gene}

    orth = pd.read_csv(DATA / "1.Databases/Ensembl_orthologues/release_116/raw/human_mouse_orthologues.tsv", sep="\t")
    one2one = orth[orth["Mouse homology type"] == "ortholog_one2one"].drop_duplicates("Gene name")
    one2one = dict(zip(one2one["Gene name"], one2one["Mouse gene name"]))
    mouse_gene, mouse_note = {}, {}
    for g in cand.gene:
        if g in MOUSE_OVERRIDE:
            mouse_gene[g], mouse_note[g] = MOUSE_OVERRIDE[g]
        elif g in one2one:
            mouse_gene[g], mouse_note[g] = one2one[g], "Ensembl 116 one2one"
        else:
            mouse_gene[g], mouse_note[g] = None, "no one2one orthologue"
    mouse_keys = {g: [m] if m else [] for g, m in mouse_gene.items()}

    # Gene universes (to tell "filtered" from "absent").
    with h5py.File(DATA / "scRNAseq/HGCA_Elmentaite2021/raw/epi_raw_counts02_v2.h5ad", "r") as f:
        var = f["var"]
        key = var.attrs.get("_index", "_index")
        hgca_u = set(v.decode() if isinstance(v, bytes) else v for v in var[key][:])
    mouse_u = set(pd.read_csv(DATA / "bulkRNAseq/GSE230581/raw/GSE230581_in_vivo_counts.tsv.gz", sep="\t",
                              usecols=[0, 1]).iloc[:, 1].astype(str))
    pelka_u = set(pd.read_csv(DATA / "scRNAseq/GSE178341/processed/v0.1/step3_epithelial_pseudobulk/"
                              "Pelka_epithelial_pseudobulk_counts.csv.gz", usecols=["symbol"]).symbol)
    joan_u = set(pd.read_csv(DATA / "scRNAseq/Joanito2022_syn26844071/processed/v0.1/step3_epithelial_pseudobulk/"
                             "Joanito_epithelial_pseudobulk_counts.csv.gz", usecols=["symbol"]).symbol)

    hgca = pd.read_csv(S2 / "human_HGCA_fetal_vs_adult_DEG.csv")
    gao = pd.read_csv(S2 / "human_Gao_validation.csv")
    mouse = pd.read_csv(S2 / "mouse_in_vivo_fetal_vs_adult_DEG.csv")
    bm = pd.read_csv(BM / "Literature_31_four_dataset_benchmark.csv")
    tcga = pd.read_csv(S3 / "TCGA_tumor_vs_normal_DEG.csv")
    tcga_u = set(tcga.symbol.dropna())
    tcga = tcga[tcga.T1_tested]
    jo = pd.read_csv(S3 / "Joanito_Malignant_vs_Normal.csv")
    jos = pd.read_csv(S3 / "Joanito_Malignant_vs_Normal_sensitivity_paired_cohorts.csv")
    pe = pd.read_csv(S3 / "Pelka_tumor_vs_normal_epithelium_DEG.csv")
    pp = pd.read_csv(S3 / "Pelka_tumor_vs_normal_prolif_control_DEG.csv")
    s3ev = pd.read_csv(S3 / "Literature_31_CRC_axis_audit.csv")
    s2audit = pd.read_csv(S2 / "Literature_31_gate_failure_audit.csv")

    bm_idx = bm.set_index("gene")

    def from_bm(cols):
        res = {}
        for g in cand.gene:
            v = [bm_idx.loc[g, c] if c and g in bm_idx.index else np.nan for c in cols]
            res[g] = v + (["measured"] if pd.notna(v[0]) else ["not measured (see notes)"])
        return res

    res = {
        "HGCA": lookup(hgca, "gene", human_keys, ["HGCA_log2FC", "HGCA_PValue", "HGCA_FDR"], hgca_u),
        "HGCA_prolif": lookup(hgca, "gene", human_keys,
                              ["HGCA_proliferative_log2FC", "HGCA_proliferative_PValue", "HGCA_proliferative_FDR"], hgca_u),
        "Gao": lookup(gao, "gene", human_keys, ["Gao_log2FC", None, None]),
        "Mouse_GSE230581": lookup(mouse, "mouse_gene", mouse_keys, ["Mouse_log2FC", "Mouse_PValue", "Mouse_FDR"], mouse_u),
        "Senger": from_bm(["Human_Senger_log2FC", "Human_Senger_PValue", "Human_Senger_FDR"]),
        "Fawkner": from_bm(["Human_Fawkner_fetal_median_detection_fraction", None, None]),
        "Pikkupeura_LN": from_bm(["Mouse_Pikkupeura_LN_log2FC", None, "Mouse_Pikkupeura_LN_FDR"]),
        "Pikkupeura_collagen": from_bm(["Mouse_Pikkupeura_collagen_log2FC", None, "Mouse_Pikkupeura_collagen_FDR"]),
        "GSE44433": from_bm(["Mouse_GSE44433_log2FC", "Mouse_GSE44433_PValue", "Mouse_GSE44433_FDR"]),
        "TCGA": lookup(tcga, "symbol", human_keys, ["T1_log2FC", "T1_PValue", "T1_FDR"], tcga_u),
        "TCGA_paired": lookup(tcga, "symbol", human_keys, ["Paired_log2FC", None, "Paired_FDR"], tcga_u),
        "TCGA_GTEx": lookup(tcga, "symbol", human_keys, ["GTEx_log2FC", None, "GTEx_FDR"], tcga_u),
        "Joanito": lookup(jo, "symbol", human_keys, ["log2FC", "PValue", "FDR"], joan_u),
        "Joanito_sens": lookup(jos, "symbol", human_keys, ["log2FC", "PValue", "FDR"], joan_u),
        "Pelka": lookup(pe, "symbol", human_keys, ["log2FC", "PValue", "FDR"], pelka_u),
        "Pelka_P": lookup(pp, "symbol", human_keys, ["log2FC", "PValue", "FDR"], pelka_u),
    }
    # Mouse-only candidates: human datasets cannot measure them.
    for key, *_ in DATASETS:
        for g in ("LY6A", "REG3B"):
            if key not in ("Mouse_GSE230581", "Pikkupeura_LN", "Pikkupeura_collagen", "GSE44433"):
                res[key][g] = [np.nan] * 3 + ["no human one2one orthologue"]
        if key in ("Mouse_GSE230581", "Pikkupeura_LN", "Pikkupeura_collagen", "GSE44433") and \
                not mouse_gene.get("SPRR1A"):
            res[key]["SPRR1A"] = [np.nan] * 3 + ["no mouse orthologue"]

    base = pd.DataFrame({
        "marker": cand.gene,
        "literature_label": cand.literature_label,
        "human_source_symbol": [to_src.get(g, g) if g not in ("LY6A", "REG3B") else "" for g in cand.gene],
        "mouse_gene_used": [mouse_gene[g] or "" for g in cand.gene],
        "mouse_mapping": [mouse_note[g] for g in cand.gene],
    })
    wide = base.copy()
    status = base[["marker"]].copy()
    for key, *_ in DATASETS:
        vals = pd.DataFrame([res[key][g] for g in cand.gene], columns=STATS + ["status"])
        for s in STATS:
            wide[f"{key}__{s}"] = pd.to_numeric(vals[s], errors="coerce")
        status[key] = vals.status
    s2map = s2audit.set_index("Gene")
    s3map = s3ev.set_index("gene")
    wide["Step2_final_(invalidated)"] = [s2map.loc[g, "Final"] if g in s2map.index else "" for g in cand.gene]
    wide["Step2_first_failed_gate"] = [s2map.loc[g, "first_failed_gate"] if g in s2map.index else "" for g in cand.gene]
    wide["Step3_final_label"] = [s3map.loc[g, "final_label"] if g in s3map.index else "" for g in cand.gene]
    wide["Step3_bulk_support"] = [s3map.loc[g, "bulk_support"] if g in s3map.index else "" for g in cand.gene]
    wide = wide.replace({np.nan: None})

    wide.to_csv(OUT / "Literature_31_marker_summary_statistics.csv", index=False)
    status.to_csv(OUT / "Literature_31_marker_measurement_status.csv", index=False)
    write_xlsx(wide, status, base, DATASETS, "Literature_31_marker_summary_statistics.xlsx", restricted=True)

    # Public version: no Joanito values and no Joanito-derived labels.
    pub_sets = [d for d in DATASETS if not d[0].startswith("Joanito")]
    drop = [c for c in wide.columns if c.startswith("Joanito") or c == "Step3_final_label"]
    pub, pub_status = wide.drop(columns=drop), status.drop(columns=[c for c in status.columns if c.startswith("Joanito")])
    pub.to_csv(OUT / "Literature_31_marker_summary_statistics_public.csv", index=False)
    pub_status.to_csv(OUT / "Literature_31_marker_measurement_status_public.csv", index=False)
    write_xlsx(pub, pub_status, base, pub_sets, "Literature_31_marker_summary_statistics_public.xlsx", restricted=False)
    print(f"{len(wide)} markers x {len(DATASETS)} dataset contrasts written to {OUT}")


def write_xlsx(wide, status, base, DATASETS, filename, restricted):
    wb = Workbook()
    ws = wb.active
    ws.title = "log2FC_P_FDR"
    arial = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    white = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {"2": "1F4E78", "2 (benchmark)": "595959", "3": "375623"}
    nbase = base.shape[1]
    # Row 1: step; row 2: dataset; row 3: contrast; row 4: statistic.
    for j, col in enumerate(base.columns, 1):
        ws.merge_cells(start_row=1, start_column=j, end_row=4, end_column=j)
        c = ws.cell(1, j, col)
        c.font, c.fill = white, PatternFill("solid", start_color="262626")
    col = nbase + 1
    for key, step, label, species, contrast, pos in DATASETS:
        for r, text in ((1, f"Step {step}"), (2, f"{label} [{species}]"), (3, f"{contrast}; + = {pos}-high")):
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 2)
            c = ws.cell(r, col, text)
            c.font, c.fill = white, PatternFill("solid", start_color=fills[step])
        for k, s in enumerate(["detection" if key == "Fawkner" else "log2FC", "P value", "FDR"]):
            c = ws.cell(4, col + k, s)
            c.font, c.fill = bold, PatternFill("solid", start_color="D9D9D9")
        col += 3
    tail = [c for c in wide.columns if c.startswith("Step")]
    for k, t in enumerate(tail):
        ws.merge_cells(start_row=1, start_column=col + k, end_row=4, end_column=col + k)
        c = ws.cell(1, col + k, t)
        c.font, c.fill = white, PatternFill("solid", start_color="7F6000")
    for i, row in enumerate(wide.itertuples(index=False), 5):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font, c.border = arial, box
            if isinstance(v, float):
                c.number_format = "0.00" if (j - nbase - 1) % 3 == 0 else "0.00E+00"
    last = ws.max_row
    for r in range(1, 5):
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            c.border = box
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 54
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    for j, w in enumerate([12, 18, 12, 11, 30], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for j in range(col, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    rule = dict(start_type="num", start_value=-3, start_color="2166AC", mid_type="num", mid_value=0,
                mid_color="FFFFFF", end_type="num", end_value=3, end_color="B2182B")
    for idx, (key, *_) in enumerate(DATASETS):
        if key == "Fawkner":
            continue
        letter = get_column_letter(nbase + 1 + 3 * idx)
        ws.conditional_formatting.add(f"{letter}5:{letter}{last}", ColorScaleRule(**rule))
    ws.freeze_panes = ws.cell(5, nbase + 1)

    # log2FC-only matrix: exactly one column per dataset.
    m = wb.create_sheet("log2FC_matrix")
    keys = [d for d in DATASETS if d[0] != "Fawkner"]
    m.append(["marker"] + [f"S{d[1]} | {d[2]}" for d in keys])
    for i, g in enumerate(wide.marker):
        m.append([g] + [wide.loc[i, f"{d[0]}__log2FC"] for d in keys])
    for c in m[1]:
        c.font, c.fill = white, PatternFill("solid", start_color="262626")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in m.iter_rows(min_row=2):
        for c in r:
            c.font, c.border = arial, box
            if isinstance(c.value, float):
                c.number_format = "0.00"
    m.row_dimensions[1].height = 60
    m.column_dimensions["A"].width = 12
    for j in range(2, m.max_column + 1):
        m.column_dimensions[get_column_letter(j)].width = 13
    m.conditional_formatting.add(f"B2:{get_column_letter(m.max_column)}{m.max_row}", ColorScaleRule(**rule))
    m.freeze_panes = "B2"

    st = wb.create_sheet("measurement_status")
    st.append(["marker"] + [f"S{d[1]} | {d[2]}" for d in DATASETS])
    for row in status.itertuples(index=False):
        st.append(list(row))
    for c in st[1]:
        c.font, c.fill = white, PatternFill("solid", start_color="262626")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in st.iter_rows(min_row=2):
        for c in r:
            c.font, c.border = arial, box
            if c.value and c.value != "measured" and c.column > 1:
                c.fill = PatternFill("solid", start_color="FCE4D6")
    st.row_dimensions[1].height = 60
    for j in range(1, st.max_column + 1):
        st.column_dimensions[get_column_letter(j)].width = 16

    nt = wb.create_sheet("Notes")
    for line in [
        "Markers: the 31 Step 1 literature-curated positive intestinal oncofetal candidates (step2_fetal/config/literature_candidates_31.tsv).",
        "Sign convention: positive log2FC = higher in fetal (Step 2) or tumour/malignant (Step 3). Blue-white-red scale spans log2FC -3..+3.",
        "P value / FDR: blank where the frozen analysis did not produce one (Gao: effect-only; Pikkupeura and TCGA paired/GTEx: only FDR retained).",
        "Fawkner-Corbett (GSE158702) has no adult arm: the value shown is the median fraction of fetal EpCAM+ cells with non-zero expression, not a fold change.",
        "Step 2 primary datasets (HGCA, Gao, GSE230581) belong to an INVALIDATED analysis (failed biological QA); values are reported for audit, not as evidence.",
        "Symbols: CCN1/CCN2 are looked up as CYR61/CTGF in annotations that predate the HGNC rename (HGCA, Gao, Senger, Fawkner, TCGA/GENCODE v26, Pelka, Joanito).",
        "Mouse: Ensembl 116 one2one orthologues; LY6A/REG3B use the mouse-defined genes Ly6a/Reg3b; CXADR uses high-confidence Cxadr (one2many). Benchmarks map by one2one Entrez only, so these are blank there.",
        "TCGA paired: 50 patients with tumour and normal; TCGA vs GTEx: source fully confounded, direction only.",
        ("Joanito (Synapse syn26844071) values are derived data under non-commercial, no-redistribution terms. Do not share this file outside the project."
         if restricted else "Public version: Joanito (Synapse syn26844071) columns and the Joanito-derived Step 3 final label are omitted under its data-use terms."),
        "Regenerate: python3 marker_summary/scripts/build_marker_summary.py",
    ]:
        nt.append([line])
        nt.cell(nt.max_row, 1).font = arial
    nt.column_dimensions["A"].width = 150
    wb.save(OUT / filename)


if __name__ == "__main__":
    main()
