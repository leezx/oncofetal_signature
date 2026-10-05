#!/usr/bin/env python3
"""Tissue and cancer breadth annotation (supplementary; NOT signature construction).

Applies the frozen plan core_oncofetal/docs/BREADTH_ANNOTATION_PLAN.md (commit 9c958ee) once.
  A. Fetal tissue breadth: Cao et al. 2020 fetal atlas (GSE156793) aggregated
     tables; organ-epithelium CPM (cell-count-weighted over prespecified
     epithelial cell types), expressed = CPM >= 10, n_expr over 8 organs.
  B. Cancer reactivation breadth: UCSC Toil TCGA + GTEx RSEM TPM; carcinomas vs
     GTEx matched tissue + TCGA solid-tissue normal; delta of medians of
     log2(TPM+1), Mann-Whitney, BH per cancer type over the annotated universe;
     up = delta >= 1 and FDR < 0.05.
Genes: 338 cross-species fetal–CRC candidates ∪ Literature-31 ∪ calibrators.
Membership labels are restricted -> outputs git-ignored, mirrored to DATA.
Outputs: results/Breadth_annotation.xlsx, results/Breadth_CIOC8_heatmap.pdf,
         results/Breadth_2D_classification.pdf
Usage (repo root): python3 core_oncofetal/scripts/breadth_annotation.py
"""
import gzip
import pathlib
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import mannwhitneyu

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA")
CAO = DATA / "scRNAseq/GSE156793_Cao2020_fetal_atlas"
TOIL = DATA / "bulkRNAseq/UCSC_Toil_TCGA_GTEx/raw"
HGNC = DATA / "1.Databases/HGNC_gene_id_mapping/raw/hgnc_custom_download.tsv"
OUT = ROOT / "core_oncofetal/results"
RESTRICTED = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal"
CPM_EXPR, MIN_EPI_CELLS, CONF = 10.0, 200, -3.0
UP_LFC, FDR, MIN_N = 1.0, 0.05, 10
ORGAN_EPI = {
    "Intestine": ["Intestinal epithelial cells"],
    "Stomach": ["Ciliated epithelial cells", "Goblet cells", "MUC13_DMBT1 positive cells", "Neuroendocrine cells",
                "Parietal and chief cells", "Squamous epithelial cells"],
    "Pancreas": ["Acinar cells", "Ductal cells", "Islet endocrine cells"],
    "Kidney": ["Metanephric cells", "Ureteric bud cells"],
    "Lung": ["Bronchiolar and alveolar epithelial cells", "Ciliated epithelial cells", "Neuroendocrine cells",
             "Squamous epithelial cells"],
    "Liver": ["Hepatoblasts"],
    "Placenta": ["Extravillous trophoblasts", "Syncytiotrophoblasts and villous cytotrophoblasts",
                 "Trophoblast giant cells"],
    "Thymus": ["Thymic epithelial cells"],
}
CELL_ASSAY = ["Intestine", "Stomach", "Pancreas", "Kidney"]
CANCERS = {  # TCGA detailed_category -> (abbrev, GTEx detailed categories)
    "Colon Adenocarcinoma": ("COAD", ["Colon - Transverse"]),
    "Rectum Adenocarcinoma": ("READ", ["Colon - Transverse"]),
    "Stomach Adenocarcinoma": ("STAD", ["Stomach"]),
    "Esophageal Carcinoma": ("ESCA", ["Esophagus - Mucosa"]),
    "Pancreatic Adenocarcinoma": ("PAAD", ["Pancreas"]),
    "Liver Hepatocellular Carcinoma": ("LIHC", ["Liver"]),
    "Cholangiocarcinoma": ("CHOL", ["Liver"]),
    "Lung Adenocarcinoma": ("LUAD", ["Lung"]),
    "Lung Squamous Cell Carcinoma": ("LUSC", ["Lung"]),
    "Breast Invasive Carcinoma": ("BRCA", ["Breast - Mammary Tissue"]),
    "Kidney Clear Cell Carcinoma": ("KIRC", ["Kidney - Cortex"]),
    "Kidney Papillary Cell Carcinoma": ("KIRP", ["Kidney - Cortex"]),
    "Kidney Chromophobe": ("KICH", ["Kidney - Cortex"]),
    "Bladder Urothelial Carcinoma": ("BLCA", ["Bladder"]),
    "Prostate Adenocarcinoma": ("PRAD", ["Prostate"]),
    "Thyroid Carcinoma": ("THCA", ["Thyroid"]),
    "Uterine Corpus Endometrioid Carcinoma": ("UCEC", ["Uterus"]),
    "Cervical & Endocervical Cancer": ("CESC", ["Cervix - Ectocervix", "Cervix - Endocervix"]),
    "Ovarian Serous Cystadenocarcinoma": ("OV", ["Ovary"]),
    "Adrenocortical Cancer": ("ACC", ["Adrenal Gland"]),
    "Head & Neck Squamous Cell Carcinoma": ("HNSC", []),
}
CRC = ("COAD", "READ")
CALIB = {"EPCAM": "pan-epithelial", "KRT8": "pan-epithelial", "CDH1": "pan-epithelial",
         "CDX2": "intestine-biased", "VIL1": "intestine-biased", "CDH17": "intestine-biased",
         "PTPRC": "non-epithelial", "COL1A2": "non-epithelial", "MKI67": "pan-cancer up", "TOP2A": "pan-cancer up"}


def hgnc():
    h = pd.read_csv(HGNC, sep="\t", dtype=str, usecols=["Approved symbol", "Previous symbols", "Ensembl gene ID"])
    ens = dict(zip(h["Ensembl gene ID"].dropna(), h.loc[h["Ensembl gene ID"].notna(), "Approved symbol"]))
    prev = {}
    for a, p in zip(h["Approved symbol"], h["Previous symbols"].fillna("")):
        for s in p.split(","):
            if s.strip():
                prev.setdefault(s.strip(), a)
    return ens, prev


def gene_sets():
    core = pd.read_csv(OUT / "Core_oncofetal_gate_statistics.csv")
    cioc = core.loc[core["Core label (v4.0)"] == "CIOC", "marker"].tolist()
    l2 = pd.read_csv(OUT / "Genomewide_fetal_CRC_candidates_calls.csv")
    c338 = l2.loc[l2["Cross-species fetal–CRC candidate"] == "YES", "Gene"].tolist()
    ext = pd.read_csv(OUT / "Extended_CIOC_calls.csv")
    e222 = ext.loc[ext["Extended CIOC"] == "YES", "Gene"].tolist()
    ec = pd.read_csv(OUT / "ECOS_38_calls.csv")
    e38 = ec.loc[ec["ECOS-38"] == "YES", "Gene"].tolist()
    lit = pd.read_csv(ROOT / "step2_fetal/config/literature_candidates_31.tsv", sep="\t").gene.tolist()
    assert (len(cioc), len(c338), len(e222), len(e38), len(lit)) == (8, 338, 222, 38, 31)
    return cioc, c338, e222, e38, lit


# ---------------------------------------------------------------- A. Cao fetal breadth
def fetal(genes, ens, prev):
    s6 = pd.read_csv(CAO / "raw/GSE156793_S6_gene_expression_celltype.txt.gz", index_col=0)
    counts = pd.read_csv(CAO / "Cao2020_celltype_counts.csv", index_col=0).n_cells
    s2 = pd.read_csv(CAO / "raw/GSE156793_S2_Metadata_genes.txt.gz")
    s2["ens"] = s2.gene_id.str.split(".").str[0]
    s2["sym"] = [ens.get(e) or prev.get(s, s) for e, s in zip(s2.ens, s2.gene_short_name)]
    s6.index = s6.index.map(dict(zip(s2.gene_id, s2.sym)))
    s6 = s6[~s6.index.duplicated(keep="first")]
    org = {}
    for o, types in ORGAN_EPI.items():
        cols = [f"{o}-{t}" for t in types]
        w = counts.reindex(cols).to_numpy(float)
        assert w.sum() >= MIN_EPI_CELLS, o
        org[o] = s6[cols].to_numpy() @ (w / w.sum())
    E = pd.DataFrame(org, index=s6.index).reindex(genes)
    non_epi = [c for c in s6.columns if c.startswith("Intestine-") and c != "Intestine-Intestinal epithelial cells"]
    nonmax = s6[non_epi].reindex(genes)
    rows = []
    for g in genes:
        e = E.loc[g]
        if e.isna().all():
            rows.append(dict(gene=g, fetal_class="not measured in Cao"))
            continue
        expr = e >= CPM_EXPR
        n = int(expr.sum())
        if not expr["Intestine"]:
            cls = "not detected in fetal intestinal epithelium"
        elif n >= 5:
            cls = "pan-fetal epithelial"
        elif n >= 3:
            cls = "multi-organ fetal epithelial"
        else:
            cls = "intestine-biased fetal"
        x = np.log2(e.to_numpy() + 1)
        tau = float(np.sum(1 - x / x.max()) / (len(x) - 1)) if x.max() > 0 else np.nan
        top = nonmax.loc[g].idxmax().replace("Intestine-", "")
        ratio = np.log2((e["Intestine"] + 1) / (nonmax.loc[g].max() + 1))
        rows.append(dict(gene=g, fetal_class=cls, n_organ_epithelia_expressed=n,
                         organs_expressed=", ".join(e.index[expr]),
                         n_cell_assay_organs_expressed=int(expr[CELL_ASSAY].sum()), tau=round(tau, 3),
                         fetal_intestine_epi_vs_top_nonepi_log2=round(ratio, 2), fetal_intestine_top_nonepi=top,
                         fetal_intestine_compartment_flag="non-epithelial dominated" if ratio < CONF else "",
                         **{f"CPM {o} epi": round(v, 2) for o, v in e.items()}))
    return pd.DataFrame(rows).set_index("gene")


# ---------------------------------------------------------------- B. Toil pan-cancer breadth
def cancer(genes, ens, prev):
    ph = pd.read_csv(TOIL / "TcgaTargetGTEX_phenotype.txt.gz", sep="\t", encoding="latin1").set_index("sample")
    want = set(genes)
    rows, ids = [], []
    with gzip.open(TOIL / "TcgaTargetGtex_rsem_gene_tpm.gz", "rt") as fh:
        header = fh.readline().rstrip("\n").split("\t")[1:]
        for line in fh:
            gid, rest = line.split("\t", 1)
            sym = ens.get(gid.split(".")[0])
            if sym in want:
                ids.append(sym)
                rows.append(np.array(rest.split("\t"), dtype=np.float32))
    X = pd.DataFrame(np.vstack(rows), index=ids, columns=header)
    X = X[~X.index.duplicated(keep="first")]
    X = np.log2(np.clip(2 ** X - 0.001, 0, None) + 1)  # log2(TPM+0.001) -> log2(TPM+1)
    ph = ph.reindex(X.columns)
    out, info = {}, []
    for cat, (ab, gtex) in CANCERS.items():
        tum = ph.index[(ph._study == "TCGA") & (ph.detailed_category == cat) & (ph._sample_type == "Primary Tumor")]
        nrm = ph.index[((ph._study == "TCGA") & (ph.detailed_category == cat) & (ph._sample_type == "Solid Tissue Normal"))
                       | ((ph._study == "GTEX") & ph.detailed_category.isin(gtex))]
        eligible = len(tum) >= MIN_N and len(nrm) >= MIN_N
        info.append(dict(cancer=ab, TCGA_category=cat, GTEx_reference="; ".join(gtex) or "none (TCGA normals only)",
                         n_tumour=len(tum), n_normal_GTEx=int(ph.loc[nrm, "_study"].eq("GTEX").sum()),
                         n_normal_TCGA=int(ph.loc[nrm, "_study"].eq("TCGA").sum()), eligible=eligible))
        if not eligible:
            continue
        t, n = X[tum].to_numpy(), X[nrm].to_numpy()
        delta = np.median(t, axis=1) - np.median(n, axis=1)
        p = mannwhitneyu(t, n, axis=1, alternative="two-sided").pvalue
        p = np.where(np.isnan(p), 1.0, p)
        o = np.argsort(p)
        q = np.empty_like(p)
        q[o] = np.minimum(1, np.minimum.accumulate((p[o] * len(p) / np.arange(1, len(p) + 1))[::-1])[::-1])
        out[ab] = pd.DataFrame({"delta": delta, "FDR": q}, index=X.index)
    info = pd.DataFrame(info)
    elig = [c for c in out]
    non = [c for c in elig if c not in CRC]
    rows = []
    for g in genes:
        if g not in X.index:
            rows.append(dict(gene=g, cancer_class="not measured in Toil"))
            continue
        up = {c: bool(out[c].loc[g, "delta"] >= UP_LFC and out[c].loc[g, "FDR"] < FDR) for c in elig}
        dn = {c: bool(out[c].loc[g, "delta"] <= -UP_LFC and out[c].loc[g, "FDR"] < FDR) for c in elig}
        frac = sum(up[c] for c in non) / len(non)
        crc_up = any(up[c] for c in CRC if c in up)
        if frac >= 0.5:
            cls = "pan-cancer reactivated"
        elif frac >= 0.25:
            cls = "multi-cancer reactivated"
        elif crc_up:
            cls = "CRC-biased reactivated"
        else:
            cls = "not CRC-up in bulk"
        rows.append(dict(gene=g, cancer_class=cls, CRC_up=", ".join(c for c in CRC if up.get(c)) or "",
                         n_nonCRC_up=sum(up[c] for c in non), n_nonCRC_eligible=len(non), frac_nonCRC_up=round(frac, 3),
                         nonCRC_up_types=", ".join(c for c in non if up[c]),
                         n_down=sum(dn.values()), down_types=", ".join(c for c in elig if dn[c]),
                         **{f"delta {c}": round(float(out[c].loc[g, "delta"]), 2) for c in elig},
                         **{f"FDR {c}": float(out[c].loc[g, "FDR"]) for c in elig}))
    return pd.DataFrame(rows).set_index("gene"), info, elig


def combined(f, c):
    pf = f == "pan-fetal epithelial"
    pc = c == "pan-cancer reactivated"
    if pf and pc:
        return "pan-tissue pan-cancer oncofetal"
    fl = {"pan-fetal epithelial": "pan-fetal", "multi-organ fetal epithelial": "multi-organ fetal",
          "intestine-biased fetal": "intestine-biased fetal",
          "not detected in fetal intestinal epithelium": "fetal intestinal epithelium not detected"}.get(f, f)
    cl = {"pan-cancer reactivated": "pan-cancer", "multi-cancer reactivated": "multi-cancer",
          "CRC-biased reactivated": "CRC-biased", "not CRC-up in bulk": "not CRC-up in bulk"}.get(c, c)
    return f"{fl} / {cl}"


def heatmap(d, cioc, organs, elig, path):
    sub = d.loc[cioc]
    A = np.log2(sub[[f"CPM {o} epi" for o in organs]].to_numpy(float) + 1)
    B = sub[[f"delta {c}" for c in elig]].to_numpy(float)
    U = (sub[[f"delta {c}" for c in elig]].to_numpy(float) >= UP_LFC) & (sub[[f"FDR {c}" for c in elig]].to_numpy(float) < FDR)
    fig, ax = plt.subplots(1, 2, figsize=(3 + 0.42 * (len(organs) + len(elig)), 0.45 * len(cioc) + 2.4),
                           gridspec_kw={"width_ratios": [len(organs), len(elig)]})
    im0 = ax[0].imshow(A, cmap="Blues", aspect="auto", vmin=0)
    ax[0].set_xticks(range(len(organs)), organs, rotation=60, ha="right", fontsize=8)
    ax[0].set_yticks(range(len(cioc)), cioc, fontsize=9)
    ax[0].set_title("Fetal organ epithelia (Cao 2020)\nlog2(CPM + 1)", fontsize=9)
    for i in range(A.shape[0]):
        for j in range(A.shape[1]):
            if sub[f"CPM {organs[j]} epi"].iloc[i] >= CPM_EXPR:
                ax[0].text(j, i, "•", ha="center", va="center", fontsize=8,
                           color="white" if A[i, j] > 0.6 * np.nanmax(A) else "#333333")
    plt.colorbar(im0, ax=ax[0], fraction=0.05, pad=0.02)
    lim = max(2, np.nanmax(np.abs(B)))
    im1 = ax[1].imshow(B, cmap="coolwarm", aspect="auto", vmin=-lim, vmax=lim)
    ax[1].set_xticks(range(len(elig)), elig, rotation=60, ha="right", fontsize=8)
    ax[1].set_yticks(range(len(cioc)), [""] * len(cioc))
    ax[1].set_title("Carcinoma vs matched normal (TCGA + GTEx)\nΔ median log2(TPM + 1)", fontsize=9)
    for i in range(B.shape[0]):
        for j in range(B.shape[1]):
            if U[i, j]:
                ax[1].text(j, i, "*", ha="center", va="center", fontsize=9, color="#222222")
    plt.colorbar(im1, ax=ax[1], fraction=0.03, pad=0.02)
    for k, c in enumerate(elig):
        if c in CRC:
            ax[1].get_xticklabels()[k].set_fontweight("bold")
    fig.text(0.01, 0.01, "• expressed (CPM ≥ 10); * up (Δ ≥ 1, FDR < 0.05). Cao organs: intestine/stomach/pancreas/kidney "
             "whole-cell, others nuclei. RESTRICTED.", fontsize=7, color="#555555")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path)
    plt.close(fig)


def scatter2d(d, cioc, e38, e222, path):
    x = d["n_organ_epithelia_expressed"].astype(float)
    y = d["frac_nonCRC_up"].astype(float)
    rng = np.random.default_rng(0)
    jx = x + rng.uniform(-0.25, 0.25, len(x))
    groups = [("Other cross-species candidates", ~d.index.isin(e222), "#BDBDBD"),
              ("Extended CIOC only", d.index.isin(e222) & ~d.index.isin(e38) & ~d.index.isin(cioc), "#7FA7D9"),
              ("ECOS-38", d.index.isin(e38) & ~d.index.isin(cioc), "#E08A3C"),
              ("CIOC Core", d.index.isin(cioc), "#1A1A1A")]
    fig, ax = plt.subplots(figsize=(6.4, 5))
    nd = (d["fetal_class"] == "not detected in fetal intestinal epithelium").to_numpy()
    for lab, m, col in groups:
        s = 22 if lab != "CIOC Core" else 40
        z = 3 if lab == "CIOC Core" else 2
        ax.scatter(jx[m & ~nd], y[m & ~nd], s=s, c=col, label=f"{lab} (n={int(m.sum())})",
                   edgecolors="white", linewidths=0.6, zorder=z)
        ax.scatter(jx[m & nd], y[m & nd], s=s, facecolors="none", edgecolors=col, linewidths=1.2, zorder=z)
    ax.scatter([], [], s=22, facecolors="none", edgecolors="#555555", label="hollow: not detected in fetal intestinal epithelium")
    for g in cioc:
        if pd.notna(y.get(g)):
            ax.annotate(g, (jx[g], y[g]), xytext=(4, 3), textcoords="offset points", fontsize=8)
    ax.axvline(4.5, color="#999999", lw=0.8, ls="--")
    ax.axvline(2.5, color="#999999", lw=0.8, ls=":")
    ax.axhline(0.5, color="#999999", lw=0.8, ls="--")
    ax.axhline(0.25, color="#999999", lw=0.8, ls=":")
    ax.set_xlabel("Fetal tissue breadth: organ epithelia expressed (of 8; Cao 2020)")
    ax.set_ylabel("Cancer breadth: fraction of non-CRC carcinomas up (TCGA + GTEx)")
    ax.set_xticks(range(0, 9))
    ax.set_ylim(-0.03, 1.03)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(alpha=0.15)
    ax.legend(fontsize=7, frameon=False, loc="upper left")
    ax.set_title("Fetal tissue breadth vs cancer reactivation breadth (cross-species candidates)", fontsize=9)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def main():
    cioc, c338, e222, e38, lit = gene_sets()
    ens, prev = hgnc()
    genes = list(dict.fromkeys(c338 + lit + list(CALIB)))
    F = fetal(genes, ens, prev)
    C, info, elig = cancer(genes, ens, prev)
    d = F.join(C, how="outer")
    d["combined_class"] = [combined(f, c) for f, c in zip(d.fetal_class.fillna(""), d.cancer_class.fillna(""))]
    d.insert(0, "Literature-31", ["YES" if g in lit else "" for g in d.index])
    d.insert(0, "ECOS-38", ["YES" if g in e38 else "" for g in d.index])
    d.insert(0, "Extended CIOC", ["YES" if g in e222 else "" for g in d.index])
    d.insert(0, "Cross-species 338", ["YES" if g in c338 else "" for g in d.index])
    d.insert(0, "CIOC", ["YES" if g in cioc else "" for g in d.index])
    d.insert(0, "Calibrator", [CALIB.get(g, "") for g in d.index])
    d.index.name = "Gene"
    organs = list(ORGAN_EPI)
    heatmap(d, cioc, organs, elig, OUT / "Breadth_CIOC8_heatmap.pdf")
    scatter2d(d.loc[c338], cioc, e38, e222, OUT / "Breadth_2D_classification.pdf")

    def tab(names):
        x = d.loc[names].groupby("combined_class").size().sort_values(ascending=False)
        return x
    summ = []
    for nm, s in (("CIOC 8", cioc), ("ECOS-38", e38), ("Extended CIOC 222", e222), ("Cross-species 338", c338)):
        for k, v in d.loc[s].fetal_class.value_counts().items():
            summ.append(dict(set=nm, axis="fetal", cls=k, n=v))
        for k, v in d.loc[s].cancer_class.value_counts().items():
            summ.append(dict(set=nm, axis="cancer", cls=k, n=v))
        for k, v in tab(s).items():
            summ.append(dict(set=nm, axis="combined", cls=k, n=v))
    summ = pd.DataFrame(summ)
    legend = pd.DataFrame([
        ("Scope", "Supplementary annotation only (frozen plan BREADTH_ANNOTATION_PLAN.md, 9c958ee); no gate, no membership change"),
        ("Fetal breadth", "Cao 2020 GSE156793 aggregated organ × cell-type CPM; organ epithelium = cell-weighted mean over prespecified epithelial types (8 organs ≥200 epithelial cells); expressed = CPM ≥ 10"),
        ("Fetal classes", "intestine CPM <10: not detected in fetal intestinal epithelium; else n_expr ≥5 pan-fetal epithelial; 3–4 multi-organ; ≤2 intestine-biased"),
        ("Fetal compartment flag", "log2((fetal intestinal epithelial CPM+1)/(max intestinal non-epithelial cell type CPM+1)) < −3"),
        ("Fetal caveats", "aggregated across fetuses (no donor replication); intestine/stomach/pancreas/kidney whole-cell, lung/liver/placenta/thymus nuclei (assay-confounded; cells-only breadth reported); stomach/pancreas/thymus from 3–4 fetuses"),
        ("Cancer breadth", "Toil TCGA+GTEx RSEM TPM; carcinoma primary tumours vs GTEx matched tissue + TCGA solid-tissue normal; Δ = median log2(TPM+1) difference; Mann–Whitney; BH per cancer type over annotated genes; up = Δ ≥1 & FDR <0.05"),
        ("Cancer classes", "frac of eligible non-CRC carcinomas up ≥0.5 pan-cancer; 0.25–0.5 multi-cancer; else CRC-up (COAD or READ) → CRC-biased; else not CRC-up in bulk"),
        ("Cancer caveats", "bulk composition; GTEx/TCGA batch; reference tissue ≠ cell of origin for some types (OV vs ovary; CHOL vs liver)"),
        ("Combined", "'pan-tissue pan-cancer oncofetal' only when pan-fetal epithelial AND pan-cancer reactivated"),
        ("Restriction", "Membership columns are Joanito-derived: do not share or commit"),
    ], columns=["Item", "Definition"])
    first = ["Calibrator", "CIOC", "Cross-species 338", "Extended CIOC", "ECOS-38", "Literature-31", "combined_class",
             "fetal_class", "cancer_class"]
    d = d[first + [c for c in d.columns if c not in first]].reset_index()
    sheets = [("CIOC_8", d[d.CIOC == "YES"]), ("ECOS_38", d[d["ECOS-38"] == "YES"]),
              ("Extended_222", d[d["Extended CIOC"] == "YES"]), ("Candidates_338", d[d["Cross-species 338"] == "YES"]),
              ("Literature_31", d[d["Literature-31"] == "YES"]), ("Calibrators", d[d.Calibrator != ""]),
              ("Class_summary", summ), ("Cancer_types", info), ("Legend", legend)]
    wb = Workbook()
    wb.remove(wb.active)
    for nm, df in sheets:
        ws = wb.create_sheet(nm)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            x.fill = PatternFill("solid", start_color="1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                v = None if (isinstance(v, float) and np.isnan(v)) else (bool(v) if isinstance(v, np.bool_) else v)
                x = ws.cell(i, j, v)
                x.font = Font(name="Arial", size=9, bold=(j == 1))
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = 120 if (nm == "Legend" and j == 2) else (
                30 if c in ("combined_class", "fetal_class", "cancer_class", "organs_expressed", "nonCRC_up_types") else 12)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 40
        ws.auto_filter.ref = ws.dimensions
    xl = OUT / "Breadth_annotation.xlsx"
    wb.save(xl)
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, OUT / "Breadth_CIOC8_heatmap.pdf", OUT / "Breadth_2D_classification.pdf"):
        shutil.copy2(f, RESTRICTED / f.name)
    pd.set_option("display.width", 250)
    print(info.to_string(index=False))
    print(d[d.Calibrator != ""][["Gene", "Calibrator", "fetal_class", "n_organ_epithelia_expressed", "cancer_class",
                                 "n_nonCRC_up", "CRC_up"]].to_string(index=False))
    print(d[d.CIOC == "YES"][["Gene", "fetal_class", "organs_expressed", "fetal_intestine_compartment_flag",
                              "cancer_class", "CRC_up", "n_nonCRC_up", "combined_class"]].to_string(index=False))
    print(summ[summ.axis == "combined"].to_string(index=False))


if __name__ == "__main__":
    main()
