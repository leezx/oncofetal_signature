#!/usr/bin/env python3
"""CIOC-extended: the frozen CIOC gates (H, M, C; method v4.0 + data-QC
amendment) applied genome-wide, i.e. without the Literature-31 candidate
restriction. RESTRICTED output (Gate C uses Joanito).

Universe: every human gene tested in at least one gate contrast.
Values: genome-wide gate contrasts from the panel-aware runs
(core_oncofetal/scripts/panel_contrasts.py; the genome-wide genes are the
standard filterByExpr set, BH over all tested genes; the 31 Literature
candidates carry the panel-exemption values, so CIOC is reproduced exactly),
plus the Gao >=9 W effect-only table.
Gates (unchanged thresholds, log2FC >= 0.5 & FDR < 0.05 = support):
  H  HGCA >=9 PCW, H-new2 >=9 PCW, Gao >=9 W: >= 2 evaluable and fetal-positive,
     >= 1 supported; NE if < 2 evaluable
  M  GSE230581 via Ensembl 116 one-to-one orthologue (Literature-31 genes use
     the frozen mapping); NE if no orthologue / not tested
  C  Joanito AND Pelka both supported; fail if an evaluable contrast lacks
     support; NE otherwise if one is not tested
Genome-wide genes absent from a contrast (filterByExpr or absent from the
source) are "not tested" (NE contribution), as in the original rules; only
the 31 prespecified candidates were exempt from the filter.
Outputs (git-ignored, mirrored to DATA restricted_joanito/core_oncofetal/):
  results/CIOC_extended_gate_calls.csv, results/CIOC_extended_gate_funnel.xlsx,
  results/CIOC_extended.gmt
Usage (repo root): python3 core_oncofetal/scripts/build_extended_core.py
"""
import pathlib
import shutil

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA")
WORK = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_core_oncofetal_panel_v0.1"
BM2 = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_step2_human_benchmark_v0.2"
ORTH = DATA / "1.Databases/Ensembl_orthologues/release_116/raw/human_mouse_orthologues.tsv"
OUT = ROOT / "core_oncofetal/results"
RESTRICTED = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal"
LFC, FDR = 0.5, 0.05
MOUSE_HIST = {"Ccn1": ["Cyr61"], "Ccn2": ["Ctgf"]}
HUMAN = {"HGCA_late": "HGCA ≥9 PCW", "Hnew2_late": "H-new2 ≥9 PCW", "Gao_late": "Gao ≥9 W"}
CRC = {"Joanito": "Joanito", "Pelka": "Pelka"}
CALL = {"pass": "PASS", "fail": "FAIL", "NE": "NOT EVALUABLE"}


def supported(v, q):
    return pd.notna(v) and v >= LFC and pd.notna(q) and q < FDR


def load_human(key, alias):
    f = BM2 / "de/GaoOriginal_late_GaoLI_vs_GSE103154.csv" if key == "Gao_late" else WORK / f"{key}_panel_DE.csv"
    d = pd.read_csv(f)
    d["symbol"] = d.symbol.astype(str).map(lambda s: alias.get(s, s))
    if "logCPM" in d:
        d = d.sort_values("logCPM", ascending=False, na_position="last")
    d = d.drop_duplicates("symbol")
    if "mean_CPM_case" not in d:
        d["mean_CPM_case"] = d["mean_CPM_ref"] = np.nan
    return d.set_index("symbol")


def main():
    al = pd.read_csv(ROOT / "step3_cancer/config/symbol_aliases.tsv", sep="\t")
    alias = dict(zip(al.source_symbol, al.current_hgnc_symbol))
    lit = pd.read_csv(ROOT / "step2_fetal/config/literature_candidates_31.tsv", sep="\t").gene.tolist()
    ms = pd.read_csv(ROOT / "marker_summary/results/Literature_31_marker_summary_statistics.csv").set_index("marker")
    panel = pd.read_csv(OUT / "Panel_gate_values_restricted.csv")
    core = pd.read_csv(OUT / "Core_oncofetal_gate_statistics.csv")
    cioc = set(core.loc[core["Core label (v4.0)"] == "CIOC", "marker"])

    hd = {k: load_human(k, alias) for k in list(HUMAN) + list(CRC)}
    md = pd.read_csv(WORK / "Mouse_GSE230581_panel_DE.csv").sort_values("logCPM", ascending=False)
    md = md.drop_duplicates("symbol").set_index("symbol")

    # human -> mouse one-to-one orthologue (Ensembl 116), by Ensembl ID then by name
    o = pd.read_csv(ORTH, sep="\t")
    o = o[(o["Mouse homology type"] == "ortholog_one2one") & o["Mouse gene name"].notna()]
    by_id = dict(zip(o["Gene stable ID"], o["Mouse gene name"]))
    by_name = o.drop_duplicates("Gene name").set_index("Gene name")["Mouse gene name"].to_dict()
    ensg = {}
    for k in ("HGCA_late", "Hnew2_late", "Pelka"):
        for s, gid in hd[k].gene_id.astype(str).items():
            if gid.startswith("ENSG"):
                ensg.setdefault(s, gid)
    pv = panel.set_index(["contrast", "gene"])

    genes = sorted(set().union(*[set(d.index) for d in hd.values()]) - {"nan"})
    rows = []
    for g in genes:
        r = {"Gene": g, "Literature candidate": "YES" if g in lit else ""}
        get = {}
        for k, d in hd.items():
            pk = "GaoOriginal_late" if k == "Gao_late" else k
            if g in lit and (pk, g) in pv.index:  # frozen panel values (incl. NA reasons)
                p = pv.loc[(pk, g)]
                get[k] = (p.log2FC, p.FDR, p.NA_reason if isinstance(p.NA_reason, str) else "",
                          p.get("low_count_flag") is True)
            elif g in d.index:
                x = d.loc[g]
                low = bool(x.mean_CPM_case < 1 and x.mean_CPM_ref < 1) if pd.notna(x.mean_CPM_case) else False
                get[k] = (x.log2FC, x.FDR, "", low)
            else:
                get[k] = (np.nan, np.nan, "not_tested", False)
        # mouse
        if g in lit:
            p = pv.loc[("Mouse_GSE230581", g)] if ("Mouse_GSE230581", g) in pv.index else None
            mg = ms.loc[g, "mouse_gene_used"] if isinstance(ms.loc[g, "mouse_gene_used"], str) else ""
            get["Mouse"] = ((p.log2FC, p.FDR, p.NA_reason if isinstance(p.NA_reason, str) else "",
                             p.get("low_count_flag") is True) if p is not None else (np.nan, np.nan, "no_1to1_orthologue", False))
        else:
            mg = by_id.get(ensg.get(g, ""), by_name.get(g, ""))
            hit = next((m for m in [mg] + MOUSE_HIST.get(mg, []) if m and m in md.index), None)
            if not mg:
                get["Mouse"] = (np.nan, np.nan, "no_1to1_orthologue", False)
            elif hit is None:
                get["Mouse"] = (np.nan, np.nan, "not_tested", False)
            else:
                x = md.loc[hit]
                get["Mouse"] = (x.log2FC, x.FDR, "", bool(x.mean_CPM_case < 1 and x.mean_CPM_ref < 1))
        r["Mouse orthologue"] = mg

        # H
        zero = lambda k: get[k][2] == "zero_expression"
        meas = [(0.0, np.nan) if zero(k) else get[k][:2] for k in HUMAN if pd.notna(get[k][0]) or zero(k)]
        if len(meas) < 2:
            h = "NE"
        else:
            h = "pass" if sum(v > 0 for v, _ in meas) >= 2 and any(supported(v, q) for v, q in meas) else "fail"
        # M
        v, q, why, _ = get["Mouse"]
        m = "fail" if why == "zero_expression" else ("NE" if pd.isna(v) else ("pass" if supported(v, q) else "fail"))
        # C
        failed = any((pd.notna(get[k][0]) and not supported(*get[k][:2])) or zero(k) for k in CRC)
        c = "fail" if failed else ("NE" if any(pd.isna(get[k][0]) for k in CRC) else "pass")

        def txt(k):
            v, q, why, low = get[k]
            if pd.isna(v):
                return f"NA ({why or 'not_tested'})"
            star = "**" if supported(v, q) else ("*" if pd.notna(q) and q < FDR else "")
            return f"{v:+.2f} (FDR {q:.2g}){star}" + (" [low count]" if low else "")

        r.update({
            "Gate H Human development": CALL[h],
            "Human evidence (log2FC, FDR)": " | ".join(f"{HUMAN[k]} {txt(k)}" for k in HUMAN),
            "Gate M Mouse in vivo": CALL[m],
            "Mouse evidence (log2FC, FDR)": f"GSE230581 {txt('Mouse')}",
            "Gate C CRC (Joanito + Pelka)": CALL[c],
            "CRC evidence (log2FC, FDR)": " | ".join(f"{CRC[k]} {txt(k)}" for k in CRC),
            "Extended core": "YES" if (h, m, c) == ("pass", "pass", "pass") else "",
            "Extended core human": "YES" if (h, c) == ("pass", "pass") else "",
            "CIOC (Literature-31)": "YES" if g in cioc else "",
        })
        for k in list(HUMAN) + list(CRC) + ["Mouse"]:
            r[f"{k}__log2FC"], r[f"{k}__FDR"] = get[k][0], get[k][1]
        rows.append(r)
    d = pd.DataFrame(rows)
    gates = [c for c in d.columns if c.startswith("Gate")]
    d["_n"] = d[gates].eq("PASS").sum(axis=1)
    # within members: order by the weakest supported CRC effect (descending), then name
    d["_crc"] = d[["Joanito__log2FC", "Pelka__log2FC"]].min(axis=1)
    d = d.sort_values(["Extended core", "Extended core human", "_n", "_crc"], ascending=False).drop(columns=["_n", "_crc"])

    ext = d[d["Extended core"] == "YES"]
    assert cioc <= set(ext.Gene), f"CIOC not reproduced: {cioc - set(ext.Gene)}"
    hp, mp, cp = (d[g].eq("PASS") for g in gates)
    funnel = pd.DataFrame([
        ("Genes tested in ≥1 gate contrast (universe)", len(d)),
        ("Gate H pass", int(hp.sum())),
        ("Gate H and M pass", int((hp & mp).sum())),
        ("Gate H, M and C pass = Extended core", int((hp & mp & cp).sum())),
        ("Extended core human (H and C, ignoring M)", int((hp & cp).sum())),
        ("Extended core ∩ Literature-31", int((hp & mp & cp & d["Literature candidate"].eq("YES")).sum())),
        ("CIOC (Literature-31) recovered in Extended core", f"{len(cioc & set(ext.Gene))}/{len(cioc)}"),
    ], columns=["Step", "Genes"])

    show = [c for c in d.columns if "__" not in c]
    d.to_csv(OUT / "CIOC_extended_gate_calls.csv", index=False)
    legend = pd.DataFrame([
        ("Scope", "Frozen CIOC gates (method v4.0 + data-QC amendment) applied genome-wide; the only change is that the Literature-31 candidate restriction is removed. Data-driven and exploratory; not the CIOC."),
        ("Universe", "All human genes tested in ≥1 gate contrast (HGCA, H-new2, Gao, Joanito, Pelka); symbols harmonised with step3 symbol_aliases.tsv"),
        ("Gate H Human development", "HGCA ≥9 PCW, H-new2 ≥9 PCW, Gao ≥9 W: ≥2 evaluable and fetal-positive and ≥1 with log2FC≥0.5 & FDR<0.05; <2 evaluable = not evaluable"),
        ("Gate M Mouse in vivo", "GSE230581 E16.5 epithelium vs adult crypt: log2FC≥0.5 & FDR<0.05 via Ensembl 116 one-to-one orthologue (Literature-31 genes: frozen mapping)"),
        ("Gate C CRC", "Joanito malignant vs normal AND Pelka tumour vs normal epithelium (epithelial pseudobulk): each log2FC≥0.5 & FDR<0.05"),
        ("FDR", "edgeR QL; BH over all genes tested in each contrast (genome-wide filterByExpr set + rescued Literature-31 genes); Gao: Welch P, BH"),
        ("NA (not_tested)", "Gene not in the contrast after genome-wide filterByExpr or absent from the source; contributes not evaluable, not fail. Only the 31 prespecified candidates were exempt from the filter."),
        ("Cells", "log2FC (FDR); ** log2FC≥0.5 & FDR<0.05; * FDR<0.05 below effect threshold; [low count] mean CPM<1 in both arms"),
        ("Extended core", "YES = passes H, M and C; Extended core human = passes H and C ignoring M; CIOC (Literature-31) = frozen 8-gene core"),
        ("Order", "Extended core first, then Extended core human; within groups by number of gates passed, then weaker CRC log2FC (min of Joanito, Pelka) descending"),
        ("Restriction", "Contains Joanito-derived results (Synapse data-use terms): do not share or commit"),
    ], columns=["Item", "Definition"])

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    fill = {"PASS": "C6EFCE", "FAIL": "F8CBAD", "NOT EVALUABLE": "D9D9D9"}
    yes_cols = ("Extended core", "Extended core human", "CIOC (Literature-31)", "Literature candidate")
    sheets = [("Extended_core", ext[show]), ("Extended_core_human", d[d["Extended core human"] == "YES"][show]),
              ("Funnel", funnel), ("All_genes", d[show]), ("Legend", legend)]
    for name, df in sheets:
        ws = wb.create_sheet(name)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            x.fill = PatternFill("solid", start_color="7030A0" if c.startswith(("Gate", "Extended", "CIOC")) else "1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        big = len(df) > 2000
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                c = df.columns[j - 1]
                x = ws.cell(i, j, v)
                if big:
                    if v in fill:
                        x.fill = PatternFill("solid", start_color=fill[v])
                    continue
                x.font = Font(name="Arial", size=9, bold=(j == 1 or (v == "YES" and c in yes_cols)))
                x.alignment = Alignment(wrap_text=True, vertical="top")
                x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if v in fill:
                    x.fill = PatternFill("solid", start_color=fill[v])
                if c in yes_cols[:3] and v == "YES":
                    x.fill = PatternFill("solid", start_color="00B050")
        widths = {"Gene": 12, "Literature candidate": 11, "Mouse orthologue": 11, "Extended core": 9,
                  "Extended core human": 10, "CIOC (Literature-31)": 10, "Step": 50, "Genes": 12, "Item": 26}
        for j, c in enumerate(df.columns, 1):
            w = widths.get(c, 15 if c.startswith("Gate") else 44)
            ws.column_dimensions[get_column_letter(j)].width = 120 if (name == "Legend" and j == 2) else w
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 40
        if big:
            ws.auto_filter.ref = ws.dimensions
    xl = OUT / "CIOC_extended_gate_funnel.xlsx"
    wb.save(xl)

    gmt = OUT / "CIOC_extended.gmt"
    gmt.write_text("\t".join(["CIOC_EXTENDED_GENOME_WIDE",
                              "Frozen CIOC gates H, M, C (method v4.0 + data-QC amendment) applied genome-wide "
                              "without the Literature-31 restriction; exploratory; HGNC symbols; RESTRICTED (Joanito-derived)"]
                             + ext.Gene.tolist()) + "\n")
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, gmt, OUT / "CIOC_extended_gate_calls.csv"):
        shutil.copy2(f, RESTRICTED / f.name)
    print(funnel.to_string(index=False))
    print("Extended core:", ", ".join(ext.Gene))


if __name__ == "__main__":
    main()
