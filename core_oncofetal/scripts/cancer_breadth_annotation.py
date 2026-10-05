#!/usr/bin/env python3
"""Pan-cancer breadth of CIOC reactivation in bulk tumours (annotation; NOT validation).

Applies the frozen breadth plan revision v2 (core_oncofetal/docs/BREADTH_ANNOTATION_PLAN.md,
commit 48cff09) once. The cancer-axis statistic is imported unchanged from the v1
exploratory script (UCSC Toil TCGA + GTEx; 21 carcinomas vs GTEx matched tissue + TCGA
solid-tissue normal; delta of median log2(TPM+1); Mann-Whitney; BH per cancer type).
Per gene (338 cross-species candidates):
  CRC epithelial evidence  Gate C (+ for all 338 by construction)
  CRC bulk                 up / down / n.s. in COAD or READ (delta >= 1 or <= -1, FDR < 0.05)
  Other carcinomas up      n of 19 eligible non-CRC carcinomas
  Composition-sensitive    Gate E E1 FAIL (existing annotation)
Interpretation: CRC bulk up -> CRC-biased (<25%) / multi-cancer (25-50%) / broad (>=50%)
reactivation; CRC bulk down -> epithelial reactivation not captured by bulk; n.s. ->
bulk cancer breadth not interpretable. ", composition-sensitive" if E1 FAIL.
Outputs (restricted; git-ignored; mirrored to DATA):
  results/Cancer_breadth_annotation.xlsx, results/Cancer_breadth_CIOC8.pdf
Usage (repo root): python3 core_oncofetal/scripts/cancer_breadth_annotation.py
"""
import importlib.util
import pathlib
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("v1", ROOT / "core_oncofetal/exploratory/breadth_fetal_cancer_v1.py")
v1 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v1)
OUT, RESTRICTED, CRC, UP, FDR = v1.OUT, v1.RESTRICTED, v1.CRC, v1.UP_LFC, v1.FDR


def interpret(bulk, frac, comp):
    if bulk == "up":
        s = ("broad reactivation across carcinomas" if frac >= 0.5 else
             "multi-cancer reactivation" if frac >= 0.25 else "CRC-biased reactivation")
    elif bulk == "down":
        s = "epithelial reactivation not captured by bulk (CRC bulk down)"
    else:
        s = "bulk cancer breadth not interpretable (CRC bulk n.s.)"
    return s + (", composition-sensitive" if comp else "")


def main():
    cioc, c338, e222, e38, _ = v1.gene_sets()
    ens, prev = v1.hgnc()
    C, info, elig = v1.cancer(c338, ens, prev)
    non = [c for c in elig if c not in CRC]
    E = pd.read_csv(OUT / "Extended_CIOC_calls.csv").set_index("Gene")
    e1 = E["E1 non-epithelial attribution (annotation)"].eq("FAIL")
    rows = []
    for g in c338:
        r = C.loc[g]
        if pd.isna(r.get("frac_nonCRC_up")):
            rows.append(dict(Gene=g, Interpretation="not measured in Toil"))
            continue
        up = [c for c in CRC if r[f"delta {c}"] >= UP and r[f"FDR {c}"] < FDR]
        dn = [c for c in CRC if r[f"delta {c}"] <= -UP and r[f"FDR {c}"] < FDR]
        bulk = "up" if up else ("down" if dn else "n.s.")
        comp = bool(e1.get(g, False))
        rows.append({
            "Gene": g,
            "CRC epithelial evidence (Gate C)": "+",
            "CRC bulk": bulk + (f" ({', '.join(up or dn)})" if (up or dn) else ""),
            "Δ COAD": r["delta COAD"], "FDR COAD": r["FDR COAD"], "Δ READ": r["delta READ"], "FDR READ": r["FDR READ"],
            "Other carcinomas up": f"{int(r.n_nonCRC_up)}/{len(non)}",
            "Other carcinomas up (types)": r.nonCRC_up_types,
            "Carcinomas down (types)": r.down_types,
            "Composition-sensitive (Gate E E1 FAIL)": "yes" if comp else "",
            "Interpretation": interpret(bulk, r.n_nonCRC_up / len(non), comp),
            **{f"Δ {c}": r[f"delta {c}"] for c in non}, **{f"FDR {c}": r[f"FDR {c}"] for c in non}})
    d = pd.DataFrame(rows).set_index("Gene")
    d.insert(0, "ECOS-38", ["YES" if g in e38 else "" for g in d.index])
    d.insert(0, "Extended CIOC", ["YES" if g in e222 else "" for g in d.index])
    d.insert(0, "CIOC", ["YES" if g in cioc else "" for g in d.index])
    d = d.reset_index()

    # figure: CIOC x carcinomas
    sub = d.set_index("Gene").loc[cioc]
    B = sub[[f"Δ {c}" for c in elig]].to_numpy(float)
    U = np.column_stack([(sub[f"Δ {c}"] >= UP) & (sub[f"FDR {c}"] < FDR) for c in elig])
    fig, ax = plt.subplots(figsize=(14, 0.48 * len(cioc) + 2.6))
    lim = max(2, np.nanmax(np.abs(B)))
    im = ax.imshow(B, cmap="coolwarm", aspect="auto", vmin=-lim, vmax=lim)
    ax.set_xticks(range(len(elig)), elig, rotation=60, ha="right", fontsize=8)
    ax.set_yticks(range(len(cioc)), cioc, fontsize=9)
    for k, c in enumerate(elig):
        if c in CRC:
            ax.get_xticklabels()[k].set_fontweight("bold")
    for i in range(B.shape[0]):
        for j in range(B.shape[1]):
            if U[i, j]:
                ax.text(j, i, "*", ha="center", va="center", fontsize=9, color="#222222")
        ax.text(len(elig) - 0.2, i, "  " + sub["Interpretation"].iloc[i], ha="left", va="center", fontsize=7.5,
                color="#333333", clip_on=False)
    ax.set_title("Pan-cancer breadth of CIOC reactivation in bulk tumours: carcinoma vs matched normal "
                 "(TCGA + GTEx), Δ median log2(TPM + 1)", fontsize=9, loc="left")
    cax = fig.add_axes([0.08, 0.13, 0.2, 0.02])
    fig.colorbar(im, cax=cax, orientation="horizontal").set_label("Δ median log2(TPM + 1)", fontsize=7)
    fig.text(0.01, 0.01, "* up (Δ ≥ 1, FDR < 0.05). Bulk annotation only; CRC epithelial reactivation is "
             "established by the epithelial gate (Joanito + Pelka). RESTRICTED.", fontsize=7, color="#555555")
    fig.subplots_adjust(left=0.08, right=0.62, bottom=0.3, top=0.88)
    fig.savefig(OUT / "Cancer_breadth_CIOC8.pdf")
    plt.close(fig)

    summ = []
    for nm, s in (("CIOC 8", cioc), ("ECOS-38", e38), ("Extended CIOC 222", e222), ("Cross-species 338", c338)):
        base = d.set_index("Gene").loc[s, "Interpretation"].str.replace(", composition-sensitive", "", regex=False)
        for k, v in base.value_counts().items():
            summ.append(dict(set=nm, interpretation=k, n=v,
                             of_which_composition_sensitive=int(d.set_index("Gene").loc[s][base.eq(k).values]
                                                                ["Composition-sensitive (Gate E E1 FAIL)"].eq("yes").sum())))
    summ = pd.DataFrame(summ)
    legend = pd.DataFrame([
        ("Title", "Pan-cancer breadth of CIOC reactivation in bulk tumours — annotation, not validation; frozen plan v2 (48cff09)"),
        ("Scope", "Cancer breadth only. Fetal tissue specificity is not part of the CIOC definition and is not reported."),
        ("CRC epithelial evidence", "Gate C: Joanito + Pelka epithelial pseudobulk (positive for all 338 by construction); bulk cannot overrule it"),
        ("Bulk data", "UCSC Toil TCGA + GTEx RSEM TPM; 21 carcinomas, primary tumour vs GTEx matched tissue + TCGA solid-tissue normal; Δ median log2(TPM+1); Mann–Whitney; BH per cancer type"),
        ("CRC bulk", "up: Δ ≥ 1 & FDR < 0.05 in COAD or READ; down: Δ ≤ −1 & FDR < 0.05 (and not up); else n.s."),
        ("Interpretation", "CRC bulk up: CRC-biased (<25% of 19 other carcinomas up) / multi-cancer (25–50%) / broad (≥50%); CRC bulk down: epithelial reactivation not captured by bulk; n.s.: bulk cancer breadth not interpretable"),
        ("Composition-sensitive", "Gate E E1 FAIL in Khaliq/Che (strong non-epithelial attribution in tumour tissue)"),
        ("Why bulk can disagree", "bulk = Σ cell fraction × expression: epithelial fraction, CAF and immune content and normal-tissue composition shift bulk independently of cell-intrinsic epithelial reactivation"),
        ("Caveats", "GTEx/TCGA batch; reference tissue ≠ cell of origin for some types (OV vs ovary, CHOL vs liver)"),
        ("Restriction", "Membership columns are Joanito-derived: do not share or commit"),
    ], columns=["Item", "Definition"])
    first = ["Gene", "CIOC", "Extended CIOC", "ECOS-38", "CRC epithelial evidence (Gate C)", "CRC bulk",
             "Other carcinomas up", "Composition-sensitive (Gate E E1 FAIL)", "Interpretation"]
    d = d[first + [c for c in d.columns if c not in first]]
    sheets = [("CIOC_8", d.set_index("Gene").loc[cioc].reset_index()), ("ECOS_38", d[d["ECOS-38"] == "YES"]),
              ("Extended_222", d[d["Extended CIOC"] == "YES"]), ("Candidates_338", d),
              ("Interpretation_summary", summ), ("Cancer_types", info), ("Legend", legend)]
    wb = Workbook()
    wb.remove(wb.active)
    for nm, df in sheets:
        ws = wb.create_sheet(nm)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            x.fill = PatternFill("solid", start_color="1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                v = None if (isinstance(v, float) and np.isnan(v)) else (bool(v) if isinstance(v, np.bool_) else v)
                if isinstance(v, float):
                    v = round(v, 4) if "FDR" not in df.columns[j - 1] else float(f"{v:.3g}")
                ws.cell(i, j, v).font = Font(name="Arial", size=9, bold=(j == 1))
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = 120 if (nm == "Legend" and j == 2) else (
                48 if c in ("Interpretation", "interpretation") else 30 if "types" in c else 13)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 40
        ws.auto_filter.ref = ws.dimensions
    xl = OUT / "Cancer_breadth_annotation.xlsx"
    wb.save(xl)
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, OUT / "Cancer_breadth_CIOC8.pdf"):
        shutil.copy2(f, RESTRICTED / f.name)
    pd.set_option("display.width", 250)
    print(d.set_index("Gene").loc[cioc, ["CRC bulk", "Other carcinomas up", "Composition-sensitive (Gate E E1 FAIL)",
                                         "Interpretation"]].to_string())
    print(summ.to_string(index=False))


if __name__ == "__main__":
    main()
