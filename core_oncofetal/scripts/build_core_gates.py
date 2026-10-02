#!/usr/bin/env python3
"""Core oncofetal gate workbook (method v1.0, docs/Method_core_oncofetal_validated.md).

Gate pass = measured, log2FC >= 0.5 in the stated direction, FDR < 0.05.
Core = G1 & G2 & G3 & G4. Supportive evidence is reported, never selecting.
Values come from the marker summary (no recomputation). The full workbook
(with Joanito G3 and Core membership) is restricted and git-ignored; the
public workbook omits all Joanito-derived columns.
Usage (repo root): python3 core_oncofetal/scripts/build_core_gates.py
"""
import pathlib
import shutil

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
MS = ROOT / "marker_summary/results"
OUT = ROOT / "core_oncofetal/results"
RESTRICTED = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA/2.PROJECTS/1.TWEAKR-oncoFetal/results/"
                          "2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal")
LFC, FDR = 0.5, 0.05

# key in marker summary, column label, role, species, contrast, restricted
GATES = [
    ("HGCA_late", "G1 Human: HGCA >=9 PCW fetal vs adult", "gate", "Human",
     "HGCA fetal epithelium >=9 PCW (10 donors) vs adult (7); donor pseudobulk edgeR QL", False),
    ("Mouse_GSE230581", "G2 Mouse: GSE230581 E16.5 vs adult crypt", "gate", "Mouse",
     "Pikkupeura in vivo E16.5 SI epithelium (3) vs adult crypt epithelium (3); edgeR QL", False),
    ("Joanito", "G3 CRC: Joanito malignant vs normal", "gate", "Human",
     "Joanito 2022 malignant vs normal epithelium; patient pseudobulk edgeR ~ cohort + group", True),
    ("Pelka", "G4 CRC replication: Pelka tumour vs normal", "gate", "Human",
     "Pelka 2021 tumour vs normal epithelium; patient pseudobulk edgeR", False),
]
SUPPORT = [
    ("Hnew2_late", "S Human: H-new2 Fawkner >=9 PCW vs Burclaff adult", "Human",
     "cross-study fetal vs adult epithelium; edgeR", False),
    ("GaoOriginal_late", "S Human: Gao LI >=9 W vs GSE103154 adult", "Human",
     "effect-only, cross-platform, adult n = 2", False),
    ("HGCA", "S Human: HGCA all fetal vs adult", "Human", "stage-unresolved reference; edgeR", False),
    ("GSE44433", "S Mouse: GSE44433 E17.5 vs 8-week LCM epithelium", "Mouse", "microarray limma; 5 vs 5", False),
    ("Pikkupeura_LN", "S Mouse culture: Pikkupeura laminin", "Mouse", "fetal enterospheres vs adult organoids; FDR only", False),
    ("Pikkupeura_collagen", "S Mouse culture: Pikkupeura collagen", "Mouse", "fetal enterospheres vs adult organoids; FDR only", False),
    ("Joanito_sens", "S CRC: Joanito cohorts with both groups", "Human", "sensitivity of G3", True),
    ("Pelka_P", "S CRC: Pelka tumour vs normal stem/TA (P)", "Human", "progenitor check; not a gate", False),
    ("TCGA", "S CRC bulk: TCGA tumour vs normal (T1)", "Human", "~ tissue + project; edgeR", False),
    ("TCGA_paired", "S CRC bulk: TCGA paired", "Human", "50 matched pairs", False),
    ("TCGA_GTEx", "S CRC bulk: TCGA vs GTEx colon", "Human", "direction only", False),
]
NA_HUMAN, NA_MOUSE = {"LY6A", "REG3B"}, {"SPRR1A"}


def gate_call(g, sp, v, q):
    if (sp == "Human" and g in NA_HUMAN) or (sp == "Mouse" and g in NA_MOUSE):
        return "not applicable"
    if pd.isna(v):
        return "not measured"
    if v >= LFC and pd.notna(q) and q < FDR:
        return "PASS"
    if v < 0 and pd.notna(q) and q < FDR:
        return "fail (opposite, FDR<0.05)"
    return "fail"


def support_call(g, sp, v, q):
    if (sp == "Human" and g in NA_HUMAN) or (sp == "Mouse" and g in NA_MOUSE):
        return "not applicable"
    if pd.isna(v):
        return "not measured"
    sig = pd.notna(q) and q < FDR
    if v > 0:
        return "supports (FDR<0.05)" if sig else "same direction, n.s."
    return "opposite (FDR<0.05)" if sig else "opposite direction, n.s."


def fmt(v, p, q, call):
    if call in ("not applicable", "not measured"):
        return call
    ps = "" if pd.isna(p) else f"; P {p:.2g}"
    qs = "" if pd.isna(q) else f"; FDR {q:.2g}"
    return f"log2FC {v:+.2f}{ps}{qs} | {call}"


def build(s, restricted):
    gates = [g for g in GATES if restricted or not g[5]]
    sup = [x for x in SUPPORT if restricted or not x[4]]
    long, wide = [], []
    for g in s.index:
        row = {"marker": g, "literature_label": s.loc[g, "literature_label"],
               "mouse_gene_used": s.loc[g, "mouse_gene_used"]}
        calls = {}
        for key, lab, role, sp, desc, _ in gates:
            v, p, q = s.loc[g, f"{key}__log2FC"], s.loc[g, f"{key}__PValue"], s.loc[g, f"{key}__FDR"]
            c = gate_call(g, sp, v, q)
            calls[key] = c
            row[lab] = fmt(v, p, q, c)
            long.append(dict(marker=g, evidence=lab, role="selection gate", contrast=desc, log2FC=v,
                             PValue=p, FDR=q, call=c, passes=c == "PASS"))
        for key, lab, sp, desc, _ in sup:
            v, p, q = s.loc[g, f"{key}__log2FC"], s.loc[g, f"{key}__PValue"], s.loc[g, f"{key}__FDR"]
            c = support_call(g, sp, v, q)
            row[lab] = fmt(v, p, q, c)
            long.append(dict(marker=g, evidence=lab, role="supportive (not selecting)", contrast=desc,
                             log2FC=v, PValue=p, FDR=q, call=c, passes=np.nan))
        dev = calls["HGCA_late"] == "PASS" and calls["Mouse_GSE230581"] == "PASS"
        row["Developmental Core (G1 & G2)"] = "yes" if dev else "no"
        if restricted:
            core = dev and calls["Joanito"] == "PASS" and calls["Pelka"] == "PASS"
            row["Conserved Intestinal Oncofetal Core (G1-G4)"] = "CORE" if core else "no"
        row["gates passed"] = sum(c == "PASS" for c in calls.values())
        row["first failed gate"] = next((lab.split(":")[0] for key, lab, *_ in gates if calls[key] != "PASS"), "none")
        wide.append(row)
    w = pd.DataFrame(wide)
    w = w.sort_values(["gates passed", "marker"], ascending=[False, True]).reset_index(drop=True)
    return w, pd.DataFrame(long), gates, sup


def definitions(gates, sup, restricted):
    rows = [dict(evidence=lab, role="selection gate", species=sp, contrast=desc,
                 rule=f"PASS = measured, log2FC >= {LFC}, FDR < {FDR}") for key, lab, role, sp, desc, _ in gates]
    rows += [dict(evidence=lab, role="supportive (not selecting)", species=sp, contrast=desc,
                  rule="supports = log2FC > 0 and FDR < 0.05; reported only") for key, lab, sp, desc, _ in sup]
    rows.append(dict(evidence="Core definition", role="", species="",
                     contrast="Literature 31 -> G1 -> G2 -> G3 -> G4" if restricted else
                     "Literature 31 -> G1 -> G2 (-> G3 Joanito, restricted) -> G4",
                     rule="Core = pass G1 & G2 & G3 & G4; Developmental Core = pass G1 & G2"))
    return pd.DataFrame(rows)


def write_xlsx(path, sheets, gate_labels, restricted):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    head = PatternFill("solid", start_color="1F4E78")
    gate_head = PatternFill("solid", start_color="7030A0")
    fills = {"PASS": "C6EFCE", "supports": "E2F0D9", "opposite": "F8CBAD", "fail": "F2F2F2",
             "not": "D9D9D9", "CORE": "00B050", "yes": "C6EFCE"}
    for name, df in sheets.items():
        ws = wb.create_sheet(name)
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(1, j, c)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            cell.fill = gate_head if c in gate_labels else head
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(r, 1):
                v = None if (isinstance(v, float) and np.isnan(v)) else v
                cell = ws.cell(i, j, v)
                cell.font = Font(name="Arial", size=9, bold=(j == 1))
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(v, str):
                    tag = v.split("| ")[-1]
                    for k, col in fills.items():
                        if tag.startswith(k) or v == k:
                            cell.fill = PatternFill("solid", start_color=col)
                            break
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = 14 if j == 1 else (30 if name != "Gate_definitions" else 45)
        ws.row_dimensions[1].height = 48
        ws.freeze_panes = "B2"
    wb.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    full_src = MS / "Literature_31_marker_summary_statistics.csv"
    s_full = pd.read_csv(full_src).set_index("marker")
    for restricted in (True, False):
        w, long, gates, sup = build(s_full, restricted)
        gate_labels = {lab for _, lab, *_ in gates}
        dev = w[w["Developmental Core (G1 & G2)"] == "yes"]
        gate_cols = ["marker", "literature_label"] + [lab for _, lab, *_ in gates] + \
                    (["Conserved Intestinal Oncofetal Core (G1-G4)"] if restricted else []) + ["gates passed"]
        sup_cols = ["marker"] + [lab for _, lab, *_ in sup]
        sheets = {
            "Core_gates": dev[gate_cols],
            "Core_supportive": dev[sup_cols],
            "All31_funnel": w[gate_cols[:-1] + ["Developmental Core (G1 & G2)", "gates passed", "first failed gate"]],
            "All31_supportive": w[sup_cols],
            "Long_numeric": long,
            "Gate_definitions": definitions(gates, sup, restricted),
        }
        suffix = "" if restricted else "_public"
        write_xlsx(OUT / f"Core_oncofetal_gate_statistics{suffix}.xlsx", sheets, gate_labels, restricted)
        w.to_csv(OUT / f"Core_oncofetal_gate_statistics{suffix}.csv", index=False)
        long.to_csv(OUT / ("Core_oncofetal_gate_long.csv" if restricted else "Core_oncofetal_gate_long_public.csv"),
                    index=False)
        if restricted:
            RESTRICTED.mkdir(parents=True, exist_ok=True)
            for f in ("Core_oncofetal_gate_statistics.xlsx", "Core_oncofetal_gate_statistics.csv",
                      "Core_oncofetal_gate_long.csv"):
                shutil.copy2(OUT / f, RESTRICTED / f)
            print(w[gate_cols[:1] + ["gates passed", "first failed gate",
                                     "Conserved Intestinal Oncofetal Core (G1-G4)"]].head(8).to_string(index=False))
    print("written:", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
