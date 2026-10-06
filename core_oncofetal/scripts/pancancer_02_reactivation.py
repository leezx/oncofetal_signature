#!/usr/bin/env python3
"""Pan-cancer epithelial reactivation annotation, step 2: statistics (frozen plan eaa522b, applied once).

Kang 2024 atlas pseudobulks (pancancer_01_extract_kang.py). CNV-inferred malignant epithelial
cells (Tumor, cnv tumor) vs normal epithelial cells (Normal, cnv normal, matched organ);
unit = patient x arm (cell-weighted pooling across samples, >= 20 cells), y = log2(CPM + 1)
(amendment A1: CPM = mean CP10k x 100; the frozen v1 log2(CP10k + 1) compressed genes < 100 CPM).
Tier 1: datasets with >= 2 units per arm, OLS y ~ arm + dataset; Tier 2 (cross-study): OLS y ~ arm;
>= 3 units per arm. Reactivation: log2FC >= 0.5 and BH FDR < 0.05 (BH per cancer type over the
gene universe). Breadth = k/n over evaluable non-CRC carcinomas (primary) and Tier-1-only (sensitivity).
Outputs (restricted; git-ignored; mirrored to DATA):
  results/Pancancer_epithelial_annotation.xlsx, results/Pancancer_epithelial_CIOC8.pdf
Usage (repo root): python3 core_oncofetal/scripts/pancancer_02_reactivation.py
"""
import pathlib
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA")
WORK = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-05_pancancer_epithelial_v0.1"
OUT = ROOT / "core_oncofetal/results"
RESTRICTED = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal"
EXCL = {"crc_HOLee_GSE132465", "crc_HOLee_GSE144735", "crc_HOLee_GSE132257", "crc_meta_GSE178318",
        "intestine_gca", "crc_pan_blueprint"}
MAP = {"CRC": ["Colon"], "STAD": ["Stomach"], "PAAD": ["Pancreas"], "HCC": ["Liver"], "CHOL": ["Bile Duct"],
       "LC": ["Lung"], "BRCA": ["Breast"], "RCC": ["Kidney"], "OV": ["Ovary", "Fallopian Tube"], "THCA": ["Thyroid"],
       "PRAD": ["Prostate"], "BLCA": ["Bladder"], "HNSC": ["Head and Neck"], "UCEC": ["Uterus"], "SSCC": ["Skin"]}
MIN_CELLS, MIN_UNITS_DS, MIN_UNITS, LFC, FDR = 20, 2, 3, 0.5, 0.05
CALIB = ["EPCAM", "KRT8", "CDH1", "MKI67", "TOP2A", "CEACAM5", "CDX2", "PTPRC", "COL1A2", "LYZ"]
META = ["Dataset", "Sample", "Patient", "Tissue", "Cancer type", "Organ_origin", "Celltype", "cnv_status", "n_cells"]


def units(df, genes):
    """patient x arm units: cell-weighted mean of sample pseudobulks."""
    df = df.copy()
    w = df.n_cells.to_numpy()[:, None]
    agg = (df[genes] * w).groupby([df.Dataset, df.Patient]).sum()
    n = df.groupby(["Dataset", "Patient"]).n_cells.sum()
    agg = agg.div(n, axis=0)
    keep = n >= MIN_CELLS
    return np.log2(agg[keep] * 100 + 1), n[keep]  # amendment A1: CPM = CP10k x 100; y = log2(CPM + 1)


def ols_arm(y, arm, dataset=None):
    """per-gene OLS coefficient for arm (malignant=1) with optional dataset fixed effects."""
    X = [np.ones(len(arm)), arm.astype(float)]
    if dataset is not None:
        for d in sorted(set(dataset))[1:]:
            X.append((dataset == d).astype(float))
    X = np.column_stack(X)
    Y = y.to_numpy()
    beta, *_ = np.linalg.lstsq(X, Y, rcond=None)
    resid = Y - X @ beta
    df = X.shape[0] - np.linalg.matrix_rank(X)
    s2 = (resid ** 2).sum(0) / df
    cov = np.linalg.pinv(X.T @ X)[1, 1]
    se = np.sqrt(s2 * cov)
    with np.errstate(divide="ignore", invalid="ignore"):
        t = beta[1] / se
    p = 2 * stats.t.sf(np.abs(t), df)
    p = np.where(np.isfinite(p), p, 1.0)
    return beta[1], p, df


def bh(p):
    p = np.asarray(p, float)
    o = np.argsort(p)
    q = np.empty_like(p)
    q[o] = np.minimum(1, np.minimum.accumulate((p[o] * len(p) / np.arange(1, len(p) + 1))[::-1])[::-1])
    return q


def gene_sets():
    core = pd.read_csv(OUT / "Core_oncofetal_gate_statistics.csv")
    cioc = core.loc[core["Core label (v4.0)"] == "CIOC", "marker"].tolist()
    l2 = pd.read_csv(OUT / "Genomewide_fetal_CRC_candidates_calls.csv")
    c338 = l2.loc[l2["Cross-species fetal–CRC candidate"] == "YES", "Gene"].tolist()
    e222 = pd.read_csv(OUT / "Extended_CIOC_calls.csv").query("`Extended CIOC` == 'YES'").Gene.tolist()
    e38 = pd.read_csv(OUT / "ECOS_38_calls.csv").query("`ECOS-38` == 'YES'").Gene.tolist()
    assert (len(cioc), len(c338), len(e222), len(e38)) == (8, 338, 222, 38)
    return cioc, c338, e222, e38


def main():
    cioc, c338, e222, e38 = gene_sets()
    pb = pd.read_csv(WORK / "Kang_pseudobulk_meanCP10k.csv.gz")
    genes = [c for c in pb.columns if c not in META]
    pb = pb[~pb.Dataset.isin(EXCL) & (pb.Celltype == "Epithelial")]
    for c in ("Dataset", "Patient"):
        pb[c] = pb[c].astype(str)
    mal = pb[(pb.Tissue == "Tumor") & (pb.cnv_status == "tumor")]
    nor = pb[(pb.Tissue == "Normal") & (pb.cnv_status == "normal")]
    res, info = {}, []
    for ct, orgs in MAP.items():
        ym, nm = units(mal[mal["Cancer type"] == ct], genes)
        yn, nn = units(nor[nor.Organ_origin.isin(orgs)], genes)
        dm, dn = ym.index.get_level_values(0), yn.index.get_level_values(0)
        both = [d for d in set(dm) if (dm == d).sum() >= MIN_UNITS_DS and (dn == d).sum() >= MIN_UNITS_DS]
        t1m, t1n = ym[dm.isin(both)], yn[dn.isin(both)]
        if len(t1m) >= MIN_UNITS and len(t1n) >= MIN_UNITS:
            tier, Ym, Yn, model = "Tier 1 (within-study)", t1m, t1n, "y ~ arm + dataset"
        elif len(ym) >= MIN_UNITS and len(yn) >= MIN_UNITS:
            tier, Ym, Yn, model = "Tier 2 (cross-study; weaker)", ym, yn, "y ~ arm"
        else:
            info.append(dict(cancer=ct, normal_organs="+".join(orgs), tier="not evaluable",
                             malignant_units=len(ym), normal_units=len(yn)))
            continue
        Y = pd.concat([Ym, Yn])
        arm = np.r_[np.ones(len(Ym)), np.zeros(len(Yn))]
        dsv = np.asarray(Y.index.get_level_values(0)) if tier.startswith("Tier 1") else None
        b, p, dof = ols_arm(Y, arm, dsv)
        r = pd.DataFrame({"log2FC": b, "P": p}, index=genes)
        r["FDR"] = bh(r.P)
        r["up"] = (r.log2FC >= LFC) & (r.FDR < FDR)
        r["down"] = (r.log2FC <= -LFC) & (r.FDR < FDR)
        res[ct] = r
        info.append(dict(cancer=ct, normal_organs="+".join(orgs), tier=tier, model=model, malignant_units=len(Ym),
                         normal_units=len(Yn), datasets=", ".join(sorted(set(Y.index.get_level_values(0)))),
                         residual_df=int(dof)))
    info = pd.DataFrame(info)
    ev = list(res)
    non = [c for c in ev if c != "CRC"]
    t1 = [c for c in non if res[c].attrs.get("tier") or info.set_index("cancer").loc[c, "tier"].startswith("Tier 1")]
    rows = []
    for g in dict.fromkeys(c338 + CALIB):
        if g not in genes:
            rows.append(dict(Gene=g, Interpretation="not measured in Kang"))
            continue
        k = [c for c in non if res[c].loc[g, "up"]]
        k1 = [c for c in t1 if res[c].loc[g, "up"]]
        crc = bool(res["CRC"].loc[g, "up"]) if "CRC" in res else None
        lab = ("CRC-biased" if (crc and len(k) <= 2) else f"reactivated in {len(k)}/{len(non)} other carcinomas") + \
              ("" if crc else " (not reactivated in independent CRC)")
        rows.append({"Gene": g, "CRC (independent, Tier 1)": "up" if crc else ("down" if res["CRC"].loc[g, "down"] else "n.s."),
                     "Breadth: other carcinomas up (primary)": f"{len(k)}/{len(non)}",
                     "Breadth: Tier 1 only": f"{len(k1)}/{len(t1)}",
                     "Other carcinomas up": ", ".join(k), "Carcinomas down": ", ".join(c for c in ev if res[c].loc[g, "down"]),
                     "Interpretation (descriptive)": lab,
                     **{f"log2FC {c}": res[c].loc[g, "log2FC"] for c in ev},
                     **{f"FDR {c}": res[c].loc[g, "FDR"] for c in ev}})
    d = pd.DataFrame(rows).set_index("Gene")
    d.insert(0, "ECOS-38", ["YES" if g in e38 else "" for g in d.index])
    d.insert(0, "Extended CIOC", ["YES" if g in e222 else "" for g in d.index])
    d.insert(0, "CIOC", ["YES" if g in cioc else "" for g in d.index])
    d.insert(0, "Calibrator", ["YES" if g in CALIB else "" for g in d.index])
    d = d.reset_index()

    # figure: CIOC 8 x evaluable carcinomas
    sub = d.set_index("Gene").loc[cioc]
    order = ["CRC"] + [c for c in non if c in t1] + [c for c in non if c not in t1]
    B = sub[[f"log2FC {c}" for c in order]].to_numpy(float)
    U = np.column_stack([(sub[f"log2FC {c}"] >= LFC) & (sub[f"FDR {c}"] < FDR) for c in order])
    fig, ax = plt.subplots(figsize=(12, 0.5 * len(cioc) + 2.8))
    lim = max(1.5, np.nanmax(np.abs(B)))
    im = ax.imshow(B, cmap="coolwarm", aspect="auto", vmin=-lim, vmax=lim)
    tiers = info.set_index("cancer").tier
    ax.set_xticks(range(len(order)), [c + ("" if tiers[c].startswith("Tier 1") else " †") for c in order],
                  rotation=60, ha="right", fontsize=8)
    ax.get_xticklabels()[0].set_fontweight("bold")
    ax.set_yticks(range(len(cioc)), cioc, fontsize=9)
    for i in range(B.shape[0]):
        for j in range(B.shape[1]):
            if U[i, j]:
                ax.text(j, i, "*", ha="center", va="center", fontsize=10, color="#222222")
        ax.text(len(order) - 0.2, i, f"  {sub['Breadth: other carcinomas up (primary)'].iloc[i]} other carcinomas",
                ha="left", va="center", fontsize=7.5, color="#333333", clip_on=False)
    ax.axvline(0.5, color="#555555", lw=0.8)
    ax.set_title("CIOC reactivation in malignant epithelium across carcinomas (Kang 2024; CNV-inferred malignant vs "
                 "normal epithelium, patient pseudobulk)", fontsize=9, loc="left")
    cax = fig.add_axes([0.08, 0.13, 0.2, 0.02])
    fig.colorbar(im, cax=cax, orientation="horizontal").set_label("log2FC (malignant − normal epithelium)", fontsize=7)
    fig.text(0.01, 0.01, "* log2FC ≥ 0.5 and FDR < 0.05. † Tier 2 (cross-study normal; weaker). CRC from an independent "
             "dataset (GSE166555). RESTRICTED.", fontsize=7, color="#555555")
    fig.subplots_adjust(left=0.08, right=0.8, bottom=0.3, top=0.88)
    fig.savefig(OUT / "Pancancer_epithelial_CIOC8.pdf")
    plt.close(fig)

    legend = pd.DataFrame([
        ("Question", "In which carcinomas is each gene reactivated in malignant epithelial cells relative to the corresponding normal epithelium? Annotation only; frozen plan eaa522b"),
        ("Data", "Kang et al. 2024 Nat Commun pan-cancer tumour–normal scRNA atlas (Zenodo 10651059); log1p(CP10k) matrices"),
        ("Independence", "Excluded construction-overlapping datasets: " + ", ".join(sorted(EXCL)) + ". CRC column = crc_GSE166555 only (Tier 1)"),
        ("Arms", "Malignant: Epithelial, cnv_status tumor (authors' inferCNVpy), Tumor tissue. Normal: Epithelial, cnv_status normal, Normal tissue of the matched organ"),
        ("Unit", "patient × arm pseudobulk (cell-weighted pooling across samples; ≥ 20 cells); y = log2(CPM + 1), CPM = mean CP10k × 100 (amendment A1; v1 log2(CP10k + 1) archived as a scale error); no cell-level tests"),
        ("Tiers", "Tier 1: datasets with ≥ 2 units per arm, OLS y ~ arm + dataset; Tier 2 (cross-study, weaker): OLS y ~ arm; ≥ 3 units per arm"),
        ("Reactivation", "log2FC ≥ 0.5 and BH FDR < 0.05 within cancer type across the gene universe (matches CIOC CRC gate)"),
        ("Breadth", "k/n evaluable non-CRC carcinomas reactivated (primary); Tier-1-only k/n (sensitivity); counts only, no pan-cancer dichotomy"),
        ("Descriptive label", "CRC-biased if CRC reactivated and ≤ 2 other carcinomas; else 'reactivated in k/n other carcinomas'"),
        ("Interpretation", "CIOC defines an intestinal oncofetal state, not CRC exclusivity: CRC-biased and broadly reactivated components are both CIOC components"),
        ("Secondary", "Bulk TCGA/GTEx breadth (Cancer_breadth_annotation.xlsx) is a composition-confounded secondary annotation"),
        ("Caveats", "log-normalised source (mean-CP10k pseudobulk, not raw counts); authors' CNV calls; lung cancers pooled (LC); Tier 2 cross-study; small n for CHOL, BLCA, STAD"),
        ("Restriction", "Membership columns are Joanito-derived: do not share or commit"),
    ], columns=["Item", "Definition"])
    first = ["Gene", "Calibrator", "CIOC", "Extended CIOC", "ECOS-38", "CRC (independent, Tier 1)",
             "Breadth: other carcinomas up (primary)", "Breadth: Tier 1 only", "Other carcinomas up",
             "Interpretation (descriptive)"]
    d = d[first + [c for c in d.columns if c not in first]]
    sheets = [("CIOC_8", d.set_index("Gene").loc[cioc].reset_index()), ("ECOS_38", d[d["ECOS-38"] == "YES"]),
              ("Extended_222", d[d["Extended CIOC"] == "YES"]), ("Candidates_338", d[d.Calibrator == ""]),
              ("Calibrators", d[d.Calibrator == "YES"]), ("Cancer_types", info), ("Legend", legend)]
    wb = Workbook()
    wb.remove(wb.active)
    for nm, df in sheets:
        ws = wb.create_sheet(nm)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            x.fill = PatternFill("solid", start_color="1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                v = None if (isinstance(v, float) and np.isnan(v)) else (bool(v) if isinstance(v, np.bool_) else v)
                if isinstance(v, float):
                    v = float(f"{v:.3g}") if "FDR" in df.columns[j - 1] else round(v, 3)
                ws.cell(i, j, v).font = Font(name="Arial", size=9, bold=(j == 1))
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = 120 if (nm == "Legend" and j == 2) else (
                40 if c in ("Interpretation (descriptive)", "Other carcinomas up", "datasets") else 12)
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 40
        ws.auto_filter.ref = ws.dimensions
    xl = OUT / "Pancancer_epithelial_annotation.xlsx"
    wb.save(xl)
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, OUT / "Pancancer_epithelial_CIOC8.pdf"):
        shutil.copy2(f, RESTRICTED / f.name)
    pd.set_option("display.width", 250)
    print(info.to_string(index=False))
    show = ["CRC (independent, Tier 1)", "Breadth: other carcinomas up (primary)", "Breadth: Tier 1 only",
            "Other carcinomas up", "Interpretation (descriptive)"]
    print(d.set_index("Gene").loc[cioc, show].to_string())
    print(d.set_index("Gene").loc[CALIB, show].to_string())


if __name__ == "__main__":
    main()
