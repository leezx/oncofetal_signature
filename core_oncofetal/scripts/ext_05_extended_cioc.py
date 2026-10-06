#!/usr/bin/env python3
"""Extended CIOC v2.0 (epithelial-only): apply the FROZEN rule
(core_oncofetal/docs/EXTENDED_CIOC_PLAN.md, revision v2.0, commit 72a98b6) once.

Universe: the frozen 338 cross-species fetal–CRC candidates.
Only tumour-derived epithelial cells are used (Khaliq 2022, Che 2021); stromal,
immune and endothelial expression is ignored (Gate E E1 is NOT a criterion).
  D  detected in >= 5% of tumour-derived epithelial cells (patient median) in
     >= 1 atlas (evaluability only)
  Coherence: the v1.0 framework unchanged (functions imported from
     ext_04_ecos38.py; same seed and null construction).
  Extended CIOC = D AND evaluable in both atlases AND T > 0 in both atlases.
Annotations (never selecting): empirical P per atlas, P < 0.05 either / both,
high-confidence CIOC-coherent (T > 0 both and P < 0.05 in >= 1), Gate E
compartment attribution (ratios, top non-epithelial compartment, E1 call),
Tabula Sapiens, ECOS-38 membership, CIOC core.
Reproduction check: genes tested in v1.0 must reproduce T and P exactly.
Outputs (restricted; git-ignored; mirrored to DATA restricted_joanito/core_oncofetal/):
  results/Extended_CIOC.xlsx, results/Extended_CIOC.gmt, results/Extended_CIOC_calls.csv
Usage (repo root): python3 core_oncofetal/scripts/ext_05_extended_cioc.py
"""
import importlib.util
import pathlib
import shutil

import anndata as ad
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("ecos", HERE / "ext_04_ecos38.py")
ecos = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ecos)
OUT, WORK, RESTRICTED, RAW, CIOC = ecos.OUT, ecos.WORK, ecos.RESTRICTED, ecos.RAW, ecos.CIOC
P_CUT = ecos.P_CUT


def main():
    calls = pd.read_csv(OUT / "Genomewide_fetal_CRC_candidates_calls.csv")
    l2 = calls.loc[calls["Cross-species fetal–CRC candidate"] == "YES", "Gene"].tolist()
    assert len(l2) == 338 and set(CIOC) <= set(l2)
    res = {n: ecos.resolver(ad.read_h5ad(p, backed="r").var_names) for n, p in RAW.items()}
    E = ecos.gate_e(l2, res)  # E2 = D; E1 retained as annotation only
    E = E.rename(columns={"E2 detectability": "D epithelial detectability",
                          "E1 non-epithelial attribution": "E1 non-epithelial attribution (annotation)",
                          "Gate E": "Gate E (annotation; ECOS-38 criterion)"})
    passD = E.loc[E["D epithelial detectability"] == "pass", "Gene"].tolist()

    coh, info = [], []
    for n in RAW:
        cand_sym = {g: res[n](g) for g in passD if res[n](g)}
        cioc_sym = {g: res[n](g) for g in CIOC if res[n](g)}
        assert len(cioc_sym) == 8
        c, _, i = ecos.coherence(n, cand_sym, cioc_sym)
        coh.append(c.set_index("Gene"))
        info.append(i)
    C = pd.concat(coh, axis=1)
    for n in RAW:
        C[f"{n} BH-FDR (annotation)"] = ecos.bh(C[f"{n} empirical P"])
    d = E.set_index("Gene").join(C, how="left")

    # reproduction check against v1.0 (ECOS-38 run) for genes tested in both
    old = pd.read_csv(OUT / "ECOS_38_calls.csv").set_index("Gene")
    both = old.index[old[f"{list(RAW)[0]} T (mean Fisher z)"].notna()]
    for n in RAW:
        for col in (f"{n} T (mean Fisher z)", f"{n} empirical P"):
            diff = (d.loc[both, col] - old.loc[both, col]).abs().max()
            assert diff < 1e-9, f"reproduction failed: {col} max diff {diff}"
    print(f"reproduction: {len(both)} v1.0-tested genes reproduce T and P exactly")

    T = [f"{n} T (mean Fisher z)" for n in RAW]
    P = [f"{n} empirical P" for n in RAW]
    ev = d[T].notna().all(axis=1)
    pos = (d[T] > 0).all(axis=1)
    dpass = d["D epithelial detectability"] == "pass"
    d["Extended CIOC"] = np.where(dpass & ev & pos, "YES", "")
    d["P<0.05 in either atlas"] = np.where(d[P].lt(P_CUT).any(axis=1), "yes", "")
    d["P<0.05 in both atlases"] = np.where(d[P].lt(P_CUT).all(axis=1), "yes", "")
    d["High-confidence CIOC-coherent"] = np.where(pos & ev & d[P].lt(P_CUT).any(axis=1), "YES", "")
    ecos38 = set(old.index[old["ECOS-38"] == "YES"])
    assert len(ecos38) == 38
    d["ECOS-38"] = ["YES" if g in ecos38 else "" for g in d.index]
    d["CIOC core"] = ["YES" if g in CIOC else "" for g in d.index]
    d["Tabula Sapiens LI log2 ratio (annotation)"] = old["Tabula Sapiens LI log2 ratio (annotation)"].reindex(d.index)
    miss = d["Tabula Sapiens LI log2 ratio (annotation)"].isna()
    if miss.any():
        d.loc[miss, "Tabula Sapiens LI log2 ratio (annotation)"] = ecos.ts_annotation(d.index[miss])
    d = d.reset_index()
    d.to_csv(OUT / "Extended_CIOC_calls.csv", index=False)
    info = pd.concat(info)
    info.to_csv(WORK / "Extended_CIOC_v2_coherence_strata.csv", index=False)

    ext = d[d["Extended CIOC"] == "YES"].copy()
    ext["_t"] = ext[T].mean(axis=1)
    ext = ext.sort_values(["CIOC core", "High-confidence CIOC-coherent", "_t"], ascending=False).drop(columns="_t")
    funnel = pd.DataFrame([
        ("Cross-species fetal–CRC candidates (frozen universe)", len(d)),
        ("D: fail epithelial detectability (<5% in both atlases)", int((~dpass).sum())),
        ("D pass", int(dpass.sum())),
        ("D pass, not evaluable in both atlases (<3 strata)", int((dpass & ~ev).sum())),
        ("D pass, evaluable, T ≤ 0 in ≥1 atlas", int((dpass & ev & ~pos).sum())),
        ("Extended CIOC (D ∧ evaluable ∧ T > 0 in both)", len(ext)),
        ("  of which high-confidence CIOC-coherent (P<0.05 in ≥1) [annotation]",
         int(ext["High-confidence CIOC-coherent"].eq("YES").sum())),
        ("  of which P<0.05 in both atlases [annotation]", int(ext["P<0.05 in both atlases"].eq("yes").sum())),
        ("  of which E1 non-epithelial attribution FAIL [annotation]",
         int(ext["E1 non-epithelial attribution (annotation)"].eq("FAIL").sum())),
        ("CIOC core genes in Extended CIOC", f"{int(ext['CIOC core'].eq('YES').sum())}/8"),
        ("ECOS-38 genes in Extended CIOC (separate branch; overlap reported only)",
         f"{int(ext['ECOS-38'].eq('YES').sum())}/38"),
        ("Coherence strata (Khaliq / Che)", " / ".join(str(int((info.dataset == n).sum())) for n in RAW)),
    ], columns=["Step", "Genes"])
    legend = pd.DataFrame([
        ("Rule", "Extended CIOC v2.0, frozen before computation (core_oncofetal/docs/EXTENDED_CIOC_PLAN.md, commit 72a98b6); applied once; no tuning, rescue or removal"),
        ("Question", "Among the 338 cross-species fetal–CRC candidates, which genes reproducibly track the CIOC state within the tumour-derived epithelial compartment? (state specificity, not lineage specificity)"),
        ("Use", "Broad epithelial oncofetal-state scoring in cell-resolved epithelial data. For bulk / unresolved spatial data use ECOS-38."),
        ("Cells", "Tumour-derived epithelial cells (epithelial cells from tumour tissue; malignancy not inferred), Khaliq 2022 and Che 2021; non-epithelial compartments ignored"),
        ("D", "Detected (≥1 UMI) in ≥5% of tumour-derived epithelial cells (patient median) in ≥1 atlas; evaluability only, not a lineage test"),
        ("Coherence", "Per sample stratum (≥200 cells): ~20-cell k-means metacells (≥10); candidate and CIOC score (leave-one-out for CIOC genes) residualised on S/G2M; within-stratum Spearman; T = mean Fisher z over strata (≥3 strata = evaluable)"),
        ("Membership", "D pass AND evaluable in both atlases AND T > 0 in both atlases"),
        ("Empirical P (annotation)", "One-sided vs 1,000 expression-bin-matched random gene sets; not a membership criterion"),
        ("High-confidence CIOC-coherent", "Annotation: T > 0 in both and empirical P < 0.05 in ≥1 atlas"),
        ("Gate E columns (annotation)", "log2((tumour-derived epithelial CPM+1)/(top non-epithelial CPM+1)) and E1 call; ECOS-38 criteria, not Extended CIOC criteria. Non-epithelial expression limits mixed-cell scoring; it is not evidence against epithelial oncofetal-state membership"),
        ("ECOS-38", "Separate branch from the 338 (compartment compatibility + stringent coherence) for mixed-cell data; overlap with Extended CIOC is reported, not definitional"),
        ("Interpretation note (recorded before computation)", "Tumour-derived epithelial metacells share a strong common axis (v1.0: all tested genes had T > 0 in both atlases); membership is therefore driven mainly by detectability, and the high-confidence flag carries the discriminating coherence information"),
        ("Restriction", "Universe derives from Joanito-gated results (Synapse data-use terms): do not share or commit"),
    ], columns=["Item", "Definition"])

    num = [c for c in d.columns if any(k in c for k in ("ratio", "detect", "Fisher", "empirical P", "FDR"))
           and c != "D epithelial detectability"]
    show = d.copy()
    show[num] = show[num].apply(pd.to_numeric, errors="coerce").round(4)
    first = ["Gene", "CIOC core", "Extended CIOC", "High-confidence CIOC-coherent", "ECOS-38",
             "D epithelial detectability", "P<0.05 in either atlas", "P<0.05 in both atlases"]
    show = show[first + [c for c in show.columns if c not in first]]
    ext_show = show.set_index("Gene").loc[ext.Gene].reset_index()
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    fill = {"PASS": "C6EFCE", "pass": "C6EFCE", "FAIL": "F8CBAD"}
    flag = ("Extended CIOC", "CIOC core", "High-confidence CIOC-coherent", "ECOS-38")
    for nm, df in (("Extended_CIOC", ext_show), ("All_338_candidates", show), ("Funnel", funnel),
                   ("Coherence_strata", info), ("Legend", legend)):
        ws = wb.create_sheet(nm)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            x.fill = PatternFill("solid", start_color="7030A0" if c in first[1:] else "1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                c = df.columns[j - 1]
                v = None if (isinstance(v, float) and np.isnan(v)) else v
                x = ws.cell(i, j, v)
                x.font = Font(name="Arial", size=9, bold=(j == 1 or v == "YES"))
                x.alignment = Alignment(wrap_text=(nm == "Legend"), vertical="top")
                x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if v in fill:
                    x.fill = PatternFill("solid", start_color=fill[v])
                if v == "YES" and c in flag:
                    x.fill = PatternFill("solid", start_color="00B050")
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = (120 if (nm == "Legend" and j == 2) else
                                                                 (60 if c == "Step" else 13))
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 45
        if nm in ("Extended_CIOC", "All_338_candidates"):
            ws.auto_filter.ref = ws.dimensions
    xl = OUT / "Extended_CIOC.xlsx"
    wb.save(xl)
    gmt = OUT / "Extended_CIOC.gmt"
    gmt.write_text("\t".join(["EXTENDED_CIOC", "Extended CIOC v2.0 (epithelial-only): 338 cross-species fetal–CRC "
                              "candidates with epithelial detectability and positive CIOC coherence in tumour-derived "
                              "epithelial cells of both Khaliq and Che (frozen rule, 72a98b6); for cell-resolved epithelial "
                              "data; HGNC symbols; RESTRICTED (Joanito-derived universe)"] + ext.Gene.tolist()) + "\n")
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, gmt, OUT / "Extended_CIOC_calls.csv"):
        shutil.copy2(f, RESTRICTED / f.name)
    print(funnel.to_string(index=False))


if __name__ == "__main__":
    main()
