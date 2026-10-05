#!/usr/bin/env python3
"""Extended CIOC, step 4: apply the FROZEN rules v1.0
(core_oncofetal/docs/EXTENDED_CIOC_PLAN.md, frozen at commit 4f2d9bd) once.

Universe: the 338 cross-species fetal–CRC candidates (Level 2a).
Gate E (Khaliq 2022, Che 2021; profile from ext_03_gate_e_profile.py):
  E1 fail if ratio < -3 in both atlases OR < -5 in either (one atlas measured:
  fail if < -3 there); E2 pass if tumour-derived epithelial detection >= 5% in
  >= 1 atlas; Gate E = E2 AND NOT E1-fail.
Coherence (Gate E passes): per sample stratum (>= 200 tumour-derived epithelial
  cells), k-means metacells (~20 cells, >= 10 per stratum); candidate and
  CIOC score residualised on [1, S, G2M]; within-stratum Spearman; T = mean
  Fisher z over strata (evaluable with >= 3 strata); 1,000 expression-bin-
  matched random sets (leave-one-out 7-gene sets for CIOC genes); one-sided
  empirical P. Pass: evaluable in both, T > 0 in both, P < 0.05 in >= 1.
  BH-FDR reported as annotation only.
Tabula Sapiens Large Intestine (adult normal, 10x, raw counts): epithelial vs
  top non-epithelial ratio as annotation only.
Symbols are resolved to each atlas via HGNC previous symbols (e.g. CCN2 = CTGF).
Outputs (restricted; git-ignored; mirrored to DATA restricted_joanito/core_oncofetal/):
  results/Extended_CIOC.xlsx, results/Extended_CIOC.gmt, results/Extended_CIOC_calls.csv
Usage (repo root): python3 core_oncofetal/scripts/ext_04_extended_cioc.py
"""
import pathlib
import shutil

import anndata as ad
import numpy as np
import pandas as pd
import scanpy as sc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import rankdata
from sklearn.cluster import KMeans

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA")
SC = DATA / "scRNAseq"
WORK = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_extended_cioc_v0.1"
RESTRICTED = DATA / "2.PROJECTS/1.TWEAKR-oncoFetal/results/2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal"
OUT = ROOT / "core_oncofetal/results"
HGNC = DATA / "1.Databases/HGNC_gene_id_mapping/raw/hgnc_custom_download.tsv"
TS = DATA / "1.Databases/TabulaSapiens/raw/TS_Large_Intestine.h5ad.zip"
TS_LOCAL = pathlib.Path("/private/tmp/claude-501/-Volumes-Stelligen-SSD-Stelligen-PR-oncofetal-signature/"
                        "f273181a-795c-4f16-90c7-f43ed52b212a/scratchpad/ts/TS_Large_Intestine.h5ad")
RAW = {"Khaliq2022": SC / "GSE200997_Khaliq2022/processed/v0.1/Khaliq2022_raw.h5ad",
       "Che2021": SC / "GSE178318_Che2021/processed/v0.1/Che2021_raw.h5ad"}
CIOC = ["TACSTD2", "CLU", "ANXA6", "SPP1", "CCN2", "EREG", "MSLN", "RBP1"]
E1_BOTH, E1_ANY, E2_DET = -3.0, -5.0, 0.05
MIN_STRATUM_CELLS, CELLS_PER_MC, MIN_MC, MIN_STRATA = 200, 20, 10, 3
N_NULL, N_BINS, P_CUT, SEED = 1000, 20, 0.05, 0
S_GENES = ("MCM5 PCNA TYMS FEN1 MCM2 MCM4 RRM1 UNG GINS2 MCM6 CDCA7 DTL PRIM1 UHRF1 MLF1IP CENPU HELLS RFC2 RPA2 "
           "NASP RAD51AP1 GMNN WDR76 SLBP CCNE2 UBR7 POLD3 MSH2 ATAD2 RAD51 RRM2 CDC45 CDC6 EXO1 TIPIN DSCC1 BLM "
           "CASP8AP2 USP1 CLSPN POLA1 CHAF1B BRIP1 E2F8").split()
G2M_GENES = ("HMGB2 CDK1 NUSAP1 UBE2C BIRC5 TPX2 TOP2A NDC80 CKS2 NUF2 CKS1B MKI67 TMPO CENPF TACC3 FAM64A PIMREG "
             "SMC4 CCNB2 CKAP2L CKAP2 AURKB BUB1 KIF11 ANP32E TUBB4B GTSE1 KIF20B HJURP CDCA3 HN1 JPT1 CDC20 TTK "
             "CDC25C KIF2C RANGAP1 NCAPD2 DLGAP5 CDCA2 CDCA8 ECT2 KIF23 HMMR AURKA PSRC1 ANLN LBR CKAP5 CENPE CTCF "
             "NEK2 G2E3 GAS2L3 CBX5 CENPA").split()


def resolver(var_names):
    """approved HGNC symbol -> symbol used in this atlas (approved, else a previous symbol)."""
    h = pd.read_csv(HGNC, sep="\t", dtype=str, usecols=["Approved symbol", "Previous symbols"])
    prev = {r["Approved symbol"]: [p.strip() for p in str(r["Previous symbols"]).split(",") if p.strip() and p != "nan"]
            for _, r in h.iterrows()}
    vs = set(var_names)
    return lambda g: g if g in vs else next((p for p in prev.get(g, []) if p in vs), None)


# ---------------------------------------------------------------- Gate E
def gate_e(cands, res):
    prof = pd.read_csv(WORK / "GateE_gene_profile.csv.gz").set_index(["dataset", "gene"])
    rows = []
    for g in cands:
        r = {"Gene": g}
        fails3, fails5, det = [], [], []
        for n in RAW:
            s = res[n](g)
            r[f"{n} symbol"] = s or ""
            if s is None or (n, s) not in prof.index:
                r[f"{n} log2 ratio"] = r[f"{n} epi detect"] = np.nan
                r[f"{n} top non-epi"] = ""
                continue
            p = prof.loc[(n, s)]
            r[f"{n} log2 ratio"], r[f"{n} epi detect"], r[f"{n} top non-epi"] = (p.log2_epi_vs_nonepi, p.epi_detect,
                                                                                 p.top_nonepi)
            fails3.append(p.log2_epi_vs_nonepi < E1_BOTH)
            fails5.append(p.log2_epi_vs_nonepi < E1_ANY)
            det.append(p.epi_detect >= E2_DET)
        e1 = (len(fails3) == 2 and all(fails3)) or any(fails5) or (len(fails3) == 1 and fails3[0])
        e2 = any(det)
        r["E1 non-epithelial attribution"] = "FAIL" if e1 else "pass"
        r["E2 detectability"] = "pass" if e2 else "FAIL"
        r["Gate E"] = "PASS" if (e2 and not e1) else "FAIL"
        rows.append(r)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- coherence
def metacells(name):
    a = sc.read_h5ad(RAW[name])
    o = pd.read_csv(WORK / f"{name}_cell_compartments.csv.gz", index_col=0)
    o = o[(o.compartment == "Epithelial") & (o.tissue != "Normal")]
    a = a[o.index].copy()
    a.layers["counts"] = a.X.copy()
    sc.pp.normalize_total(a, target_sum=1e4)
    sc.pp.log1p(a)
    sc.pp.highly_variable_genes(a, n_top_genes=2000)
    sc.tl.score_genes_cell_cycle(a, s_genes=[g for g in S_GENES if g in a.var_names],
                                 g2m_genes=[g for g in G2M_GENES if g in a.var_names], random_state=SEED)
    strata = []
    for st, idx in a.obs.groupby("samples", observed=True).indices.items():
        if len(idx) < MIN_STRATUM_CELLS:
            continue
        b = a[idx, a.var.highly_variable].copy()
        sc.pp.scale(b, max_value=10)
        sc.tl.pca(b, n_comps=20, random_state=SEED)
        k = len(idx) // CELLS_PER_MC
        lab = KMeans(n_clusters=k, n_init=10, random_state=SEED).fit_predict(b.obsm["X_pca"])
        cnt = a.layers["counts"][idx]
        agg = np.vstack([np.asarray(cnt[lab == j].sum(axis=0)).ravel() for j in range(k)])
        lib = agg.sum(axis=1, keepdims=True)
        keep = lib.ravel() > 0
        if keep.sum() < MIN_MC:
            continue
        x = np.log2(agg[keep] / lib[keep] * 1e6 + 1)
        cc = a.obs.iloc[idx][["S_score", "G2M_score"]].groupby(lab).mean().to_numpy()[keep]
        strata.append(dict(stratum=st, X=x.astype(np.float64), cc=cc, n_cells=len(idx), n_mc=int(keep.sum())))
    return strata, list(a.var_names)


def residualise(Y, D):
    beta, *_ = np.linalg.lstsq(D, Y, rcond=None)
    return Y - D @ beta


def zmean(X, cols):
    """mean z-score over columns; constant columns dropped. X: metacells x genes."""
    sub = X[:, cols]
    sd = sub.std(axis=0)
    ok = sd > 0
    if not ok.any():
        return np.full(X.shape[0], np.nan)
    z = (sub[:, ok] - sub[:, ok].mean(axis=0)) / sd[ok]
    return z.mean(axis=1)


def ranks_corr(A, B):
    """Spearman between rows of A (n x M) and rows of B (m x M) -> n x m."""
    ra = rankdata(A, axis=1)
    rb = rankdata(B, axis=1)
    ra = (ra - ra.mean(axis=1, keepdims=True))
    rb = (rb - rb.mean(axis=1, keepdims=True))
    ra /= np.linalg.norm(ra, axis=1, keepdims=True)
    rb /= np.linalg.norm(rb, axis=1, keepdims=True)
    return ra @ rb.T


def coherence(name, cand_sym, cioc_sym):
    """cand_sym: {gene: atlas symbol}; cioc_sym: {CIOC gene: atlas symbol}."""
    strata, var = metacells(name)
    gi = {g: i for i, g in enumerate(var)}
    allx = np.vstack([s["X"] for s in strata])
    mean_expr = allx.mean(axis=0)
    cioc_idx = {g: gi[s] for g, s in cioc_sym.items() if s in gi}
    pool_mask = mean_expr > 0
    pool_mask[list(cioc_idx.values())] = False
    pool = np.where(pool_mask)[0]
    edges = np.quantile(mean_expr[pool], np.linspace(0, 1, N_BINS + 1))
    pbin = np.clip(np.searchsorted(edges, mean_expr[pool], side="right") - 1, 0, N_BINS - 1)
    by_bin = {b: pool[pbin == b] for b in range(N_BINS)}

    def gbin(i):
        return int(np.clip(np.searchsorted(edges, mean_expr[i], side="right") - 1, 0, N_BINS - 1))

    rng = np.random.default_rng(SEED)
    families = {"full": list(cioc_idx.values())}
    for g, i in cioc_idx.items():
        families[f"loo_{g}"] = [j for j in cioc_idx.values() if j != i]
    nulls = {}
    for f, genes in families.items():
        sets = np.empty((N_NULL, len(genes)), dtype=int)
        for r in range(N_NULL):
            used = set()
            for c, j in enumerate(genes):
                choices = [x for x in by_bin[gbin(j)] if x not in used]
                x = rng.choice(choices)
                used.add(x)
                sets[r, c] = x
        nulls[f] = sets

    cands = [(g, gi[s]) for g, s in cand_sym.items() if s in gi]
    fam_of = {g: (f"loo_{g}" if g in cioc_idx else "full") for g, _ in cands}
    zsum = {g: 0.0 for g, _ in cands}
    zn = {g: 0 for g, _ in cands}
    znull = {g: np.zeros(N_NULL) for g, _ in cands}
    per_stratum = []
    for s in strata:
        X, D = s["X"], np.column_stack([np.ones(s["n_mc"]), s["cc"]])
        for f, genes in families.items():
            members = [g for g, _ in cands if fam_of[g] == f]
            if not members:
                continue
            score = zmean(X, genes)
            if np.isnan(score).any() or score.std() == 0:
                continue
            sres = residualise(score[:, None], D)[:, 0]
            nsc = np.column_stack([zmean(X, row) for row in nulls[f]])  # M x N_NULL
            okn = ~np.isnan(nsc).any(axis=0)
            nres = residualise(np.nan_to_num(nsc), D)
            cidx = [gi_ for g, gi_ in cands if fam_of[g] == f]
            Y = X[:, cidx]
            const = Y.std(axis=0) == 0
            yres = residualise(Y, D)
            obs = ranks_corr(yres.T, sres[None, :])[:, 0]
            nul = ranks_corr(yres.T, nres.T)  # n_c x N_NULL
            nul[:, ~okn] = np.nan
            for k, g in enumerate(members):
                if const[k]:
                    continue
                z = np.arctanh(np.clip(obs[k], -0.999999, 0.999999))
                zsum[g] += z
                zn[g] += 1
                znull[g] += np.nan_to_num(np.arctanh(np.clip(nul[k], -0.999999, 0.999999)))
                per_stratum.append(dict(dataset=name, stratum=s["stratum"], gene=g, rho=obs[k]))
    rows = []
    for g, _ in cands:
        n = zn[g]
        if n < MIN_STRATA:
            rows.append(dict(Gene=g, **{f"{name} n strata": n, f"{name} T (mean Fisher z)": np.nan,
                                        f"{name} empirical P": np.nan}))
            continue
        t, tn = zsum[g] / n, znull[g] / n
        rows.append(dict(Gene=g, **{f"{name} n strata": n, f"{name} T (mean Fisher z)": t,
                                    f"{name} empirical P": (1 + (tn >= t).sum()) / (N_NULL + 1)}))
    info = pd.DataFrame([dict(dataset=name, stratum=s["stratum"], n_cells=s["n_cells"], n_metacells=s["n_mc"])
                         for s in strata])
    return pd.DataFrame(rows), pd.DataFrame(per_stratum), info


# ---------------------------------------------------------------- Tabula Sapiens annotation
def ts_annotation(genes):
    path = TS_LOCAL
    if not path.exists():
        import zipfile
        path.parent.mkdir(parents=True, exist_ok=True)
        zipfile.ZipFile(TS).extractall(path.parent)
    a = ad.read_h5ad(path)
    a = a[a.obs.method == "10X"]
    X = a.layers["raw_counts"]
    sym = a.var.gene_symbol.astype(str).values
    res = resolver(sym)
    out = {}
    comps = [c for c in ["epithelial", "stromal", "endothelial", "immune"]
             if (a.obs.compartment == c).sum() >= 20]
    cpm = {}
    for c in comps:
        m = (a.obs.compartment == c).values
        v = np.asarray(X[m].sum(axis=0)).ravel()
        cpm[c] = v / v.sum() * 1e6
    idx = {s: i for i, s in enumerate(sym)}
    for g in genes:
        s = res(g)
        if s is None:
            out[g] = np.nan
            continue
        i = idx[s]
        non = max(cpm[c][i] for c in comps if c != "epithelial")
        out[g] = np.log2((cpm["epithelial"][i] + 1) / (non + 1))
    return pd.Series(out, name="Tabula Sapiens LI log2 ratio (annotation)")


def bh(p):
    p = np.asarray(p, float)
    ok = ~np.isnan(p)
    q = np.full_like(p, np.nan)
    n = ok.sum()
    o = np.argsort(p[ok])
    r = p[ok][o] * n / np.arange(1, n + 1)
    r = np.minimum.accumulate(r[::-1])[::-1]
    tmp = np.empty(n)
    tmp[o] = np.minimum(r, 1)
    q[ok] = tmp
    return q


def main():
    calls = pd.read_csv(OUT / "Genomewide_fetal_CRC_candidates_calls.csv")
    l2 = calls.loc[calls["Cross-species fetal–CRC candidate"] == "YES", "Gene"].tolist()
    assert len(l2) == 338 and set(CIOC) <= set(l2)
    res = {n: resolver(ad.read_h5ad(p, backed="r").var_names) for n, p in RAW.items()}
    E = gate_e(l2, res)
    passE = E.loc[E["Gate E"] == "PASS", "Gene"].tolist()

    coh, strat, info = [], [], []
    for n in RAW:
        cand_sym = {g: res[n](g) for g in passE if res[n](g)}
        cioc_sym = {g: res[n](g) for g in CIOC if res[n](g)}
        assert len(cioc_sym) == 8, f"{n}: CIOC symbols unresolved"
        c, s, i = coherence(n, cand_sym, cioc_sym)
        coh.append(c.set_index("Gene"))
        strat.append(s)
        info.append(i)
    C = pd.concat(coh, axis=1)
    for n in RAW:
        C[f"{n} BH-FDR (annotation)"] = bh(C[f"{n} empirical P"])
    d = E.set_index("Gene").join(C, how="left")
    ev = d[[f"{n} T (mean Fisher z)" for n in RAW]].notna().all(axis=1)
    pos = (d[[f"{n} T (mean Fisher z)" for n in RAW]] > 0).all(axis=1)
    sig = (d[[f"{n} empirical P" for n in RAW]] < P_CUT).any(axis=1)
    d["Coherence"] = np.where(d["Gate E"] != "PASS", "", np.where(ev & pos & sig, "PASS", "FAIL"))
    d["Extended CIOC"] = np.where((d["Gate E"] == "PASS") & (d.Coherence == "PASS"), "YES", "")
    d["CIOC core"] = ["YES" if g in CIOC else "" for g in d.index]
    d["Tabula Sapiens LI log2 ratio (annotation)"] = ts_annotation(d.index)
    d = d.reset_index()
    d.to_csv(OUT / "Extended_CIOC_calls.csv", index=False)
    pd.concat(strat).to_csv(WORK / "Extended_CIOC_coherence_per_stratum.csv.gz", index=False)
    info = pd.concat(info)
    info.to_csv(WORK / "Extended_CIOC_coherence_strata.csv", index=False)

    ext = d[d["Extended CIOC"] == "YES"].copy()
    ext["_t"] = ext[[f"{n} T (mean Fisher z)" for n in RAW]].mean(axis=1)
    ext = ext.sort_values(["CIOC core", "_t"], ascending=False).drop(columns="_t")
    funnel = pd.DataFrame([
        ("Cross-species fetal–CRC candidates (Level 2a)", len(d)),
        ("Gate E: fail E1 (non-epithelial attribution)", int((d["E1 non-epithelial attribution"] == "FAIL").sum())),
        ("Gate E: fail E2 (detectability)", int((d["E2 detectability"] == "FAIL").sum())),
        ("Gate E pass", int((d["Gate E"] == "PASS").sum())),
        ("Coherence: not evaluable in both atlases", int(((d["Gate E"] == "PASS") & ~ev).sum())),
        ("Coherence pass = Extended CIOC", len(ext)),
        ("CIOC core genes in Extended CIOC", f"{int(ext['CIOC core'].eq('YES').sum())}/8"),
        ("Coherence strata (Khaliq / Che)", " / ".join(str(int((info.dataset == n).sum())) for n in RAW)),
    ], columns=["Step", "Genes"])

    legend = pd.DataFrame([
        ("Rules", "Frozen rules v1.0 (core_oncofetal/docs/EXTENDED_CIOC_PLAN.md, commit 4f2d9bd), applied once; no threshold revision, no rescue"),
        ("Universe", "338 cross-species fetal–CRC candidates (frozen gates H, M, C applied genome-wide)"),
        ("Atlases", "Khaliq 2022 (GSE200997) and Che 2021 (GSE178318): independent replication within the refinement stage (not external validation)"),
        ("Tumour-derived epithelial cells", "Epithelial cells from tumour tissue; malignancy not inferred (no CNV / mutation calling)"),
        ("log2 ratio", "log2((tumour-derived epithelial CPM+1)/(top non-epithelial compartment CPM+1)), patient-median pseudobulks"),
        ("E1", "FAIL if ratio < −3 in both atlases, or < −5 in either (one atlas measured: < −3 there). −3 = >8-fold non-epithelial enrichment, chosen from blinded lineage calibrators"),
        ("E2", "pass if detected in ≥5% of tumour-derived epithelial cells (patient median) in ≥1 atlas"),
        ("Coherence", "Per sample stratum (≥200 cells): ~20-cell k-means metacells (≥10); candidate and leave-one-out CIOC score residualised on S/G2M; within-stratum Spearman; T = mean Fisher z (≥3 strata); empirical one-sided P vs 1,000 expression-matched random gene sets"),
        ("Coherence pass", "evaluable in both atlases, T > 0 in both, empirical P < 0.05 in ≥1. BH-FDR = annotation only"),
        ("Tabula Sapiens", "Adult normal large intestine (10x): epithelial vs top non-epithelial log2 ratio; annotation only, never selecting"),
        ("Restriction", "Universe derives from Joanito-gated results (Synapse data-use terms): do not share or commit"),
    ], columns=["Item", "Definition"])

    num = [c for c in d.columns if any(k in c for k in ("ratio", "detect", "Fisher", "P", "FDR"))]
    show = d.copy()
    show[num] = show[num].apply(pd.to_numeric, errors="coerce").round(4)
    first = ["Gene", "CIOC core", "Extended CIOC", "Gate E", "E1 non-epithelial attribution", "E2 detectability",
             "Coherence"]
    show = show[first + [c for c in show.columns if c not in first]]
    ext_show = show.set_index("Gene").loc[ext.Gene].reset_index()
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    fill = {"PASS": "C6EFCE", "pass": "C6EFCE", "FAIL": "F8CBAD"}
    for nm, df in (("Extended_CIOC", ext_show), ("All_338_candidates", show), ("Funnel", funnel),
                   ("Coherence_strata", info), ("Legend", legend)):
        ws = wb.create_sheet(nm)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            x.fill = PatternFill("solid", start_color="7030A0" if c in first[1:] else "1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                c = df.columns[j - 1]
                v = None if (isinstance(v, float) and np.isnan(v)) else v
                x = ws.cell(i, j, v)
                x.font = Font(name="Arial", size=9, bold=(j == 1 or v == "YES"))
                x.alignment = Alignment(wrap_text=(nm == "Legend"), vertical="top")
                x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if v in fill:
                    x.fill = PatternFill("solid", start_color=fill[v])
                if v == "YES" and c in ("Extended CIOC", "CIOC core"):
                    x.fill = PatternFill("solid", start_color="00B050")
                if isinstance(v, float) and "P" in c.split()[-1:] and v < P_CUT:
                    x.font = Font(name="Arial", size=9, bold=True)
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = (120 if (nm == "Legend" and j == 2) else
                                                                 (48 if c == "Step" else 13))
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 45
        if nm == "All_338_candidates":
            ws.auto_filter.ref = ws.dimensions
    xl = OUT / "Extended_CIOC.xlsx"
    wb.save(xl)
    gmt = OUT / "Extended_CIOC.gmt"
    gmt.write_text("\t".join(["EXTENDED_CIOC", "Extended CIOC: 338 cross-species fetal–CRC candidates ∩ Gate E "
                              "(epithelial compatibility) ∩ CIOC coherence; frozen rules v1.0 (4f2d9bd); HGNC symbols; "
                              "RESTRICTED (Joanito-derived universe)"] + ext.Gene.tolist()) + "\n")
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    for f in (xl, gmt, OUT / "Extended_CIOC_calls.csv"):
        shutil.copy2(f, RESTRICTED / f.name)
    print(funnel.to_string(index=False))


if __name__ == "__main__":
    main()
