#!/usr/bin/env python3
"""31-gene gate funnel workbook (CIOC method v4.0). RESTRICTED output.

All 31 literature candidates (Literature candidate = YES; evidence class and
primary studies are annotation) pass through the three data gates:
Gate H human development, Gate M mouse in vivo, Gate C CRC (Joanito + Pelka).
Last column Core = "YES" for CIOC members, blank otherwise.
Reads the restricted full table written by build_core_gates.py; output is
git-ignored and mirrored to DATA restricted_joanito/core_oncofetal/.
Usage (repo root): python3 core_oncofetal/scripts/build_gate_funnel.py
"""
import pathlib
import shutil

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
SRC = ROOT / "core_oncofetal/results/Core_oncofetal_gate_statistics.csv"
OUT = ROOT / "core_oncofetal/results/CIOC_gate_funnel_31.xlsx"
RESTRICTED = pathlib.Path("/Volumes/Stelligen_SSD/Stelligen/DATA/2.PROJECTS/1.TWEAKR-oncoFetal/results/"
                          "2026-10-02_step3_cancer_axis_v0.1/restricted_joanito/core_oncofetal")
CALL = {"pass": "PASS", "fail": "FAIL", "NE": "NOT EVALUABLE"}
NAMES = {"H": "Human development", "M": "Mouse in vivo", "C": "CRC"}


def col(m, key):
    return next(c for c in m.columns if key in c)


def main():
    m = pd.read_csv(SRC)
    rows = []
    for _, r in m.iterrows():
        calls = {"H": r["H human developmental"], "M": r["M mouse in vivo"], "C": r["C CRC replicated"]}
        failed = [NAMES[a] for a in "HMC" if calls[a] == "fail"]
        ne = [NAMES[a] for a in "HMC" if calls[a] == "NE"]
        filt = "; ".join((["failed: " + ", ".join(failed)] if failed else []) +
                         (["not evaluable: " + ", ".join(ne)] if ne else []))
        rows.append({
            "Gene": r.marker,
            "Literature label": r.literature_label,
            "Literature candidate": "YES",
            "Literature evidence class (annotation)": r["Literature evidence class (annotation)"],
            "Primary studies (A | B | C)": r["Primary studies (A; B; C)"],
            "Gate H Human development": CALL[calls["H"]],
            "Human evidence (log2FC, FDR)": f"HGCA≥9PCW {r[col(m, 'HGCA >=9')]} | H-new2 {r[col(m, 'H-new2')]} | "
                                            f"Gao {r[col(m, 'Gao LI')]}; support from: {r['H statistical support from']}",
            "Gate M Mouse in vivo": CALL[calls["M"]],
            "Mouse evidence (log2FC, FDR)": f"GSE230581 {r[col(m, 'GSE230581')]}",
            "Gate C CRC (Joanito + Pelka)": CALL[calls["C"]],
            "CRC evidence (log2FC, FDR)": f"Joanito {r[col(m, 'Joanito malignant')]} | Pelka {r[col(m, 'Pelka tumour')]}",
            "Filtered at": filt,
            "Discordance flags (not selecting)": r["DISCORDANCE flags (reported, not selecting)"],
            "Core": "YES" if r["Core label (v4.0)"] == "CIOC" else "",
        })
    d = pd.DataFrame(rows)
    gates = [c for c in d.columns if c.startswith("Gate")]
    d["_k"], d["_n"] = (d.Core != "YES").astype(int), d[gates].eq("PASS").sum(axis=1)
    d = d.sort_values(["_k", "_n", "Gene"], ascending=[True, False, True]).drop(columns=["_k", "_n"])
    legend = pd.DataFrame([
        ("Literature candidate", "Candidate universe = 31 literature-curated intestinal fetal / regenerative-revival / oncofetal genes; all YES (definition, not a gate)"),
        ("Literature evidence class", "Annotation from full-text audit: A fetal intestine; B regeneration/revival; C CRC oncofetal; 'provenance unresolved' = no primary source identified"),
        ("Gate H Human development", "HGCA ≥9 PCW, H-new2 ≥9 PCW, Gao ≥9 W: ≥2 fetal-positive and ≥1 with log2FC≥0.5 & FDR<0.05; <2 measured = not evaluable"),
        ("Gate M Mouse in vivo", "GSE230581 E16.5 epithelium vs adult crypt: log2FC≥0.5 & FDR<0.05 (one-to-one orthologue)"),
        ("Gate C CRC", "Joanito malignant vs normal AND Pelka tumour vs normal epithelium: each log2FC≥0.5 & FDR<0.05"),
        ("Core", "YES = passes Gates H, M and C (CIOC, method v4.0); blank = not Core"),
        ("NOT EVALUABLE", "not measured / no orthologue in the required dataset; not a biological fail"),
        ("Cells", "log2FC (FDR); ** log2FC≥0.5 & FDR<0.05; * FDR<0.05 below effect threshold; NA not measured"),
        ("Restriction", "Contains Joanito-derived results (Synapse data-use terms): do not share or commit"),
    ], columns=["Item", "Definition"])
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    fill = {"PASS": "C6EFCE", "FAIL": "F8CBAD", "NOT EVALUABLE": "D9D9D9"}
    for name, df in (("Gate_funnel_31", d), ("Legend", legend)):
        ws = wb.create_sheet(name)
        for j, c in enumerate(df.columns, 1):
            x = ws.cell(1, j, c)
            x.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            x.fill = PatternFill("solid", start_color="7030A0" if c.startswith(("Gate", "Core")) else "1F4E78")
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                c = df.columns[j - 1]
                x = ws.cell(i, j, v)
                x.font = Font(name="Arial", size=9, bold=(j == 1 or v == "YES" and c == "Core"))
                x.alignment = Alignment(wrap_text=True, vertical="top")
                x.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if v in fill:
                    x.fill = PatternFill("solid", start_color=fill[v])
                if c == "Core" and v == "YES":
                    x.fill = PatternFill("solid", start_color="00B050")
                if c.startswith("Discordance") and v not in ("none", ""):
                    x.font = Font(name="Arial", size=9, color="9C0006")
                if v == "provenance unresolved":
                    x.font = Font(name="Arial", size=9, italic=True, color="9C5700")
        widths = {"Gene": 11, "Literature label": 20, "Literature candidate": 11, "Filtered at": 28, "Core": 8}
        for j, c in enumerate(df.columns, 1):
            w = widths.get(c, 15 if c.startswith("Gate") else 38)
            ws.column_dimensions[get_column_letter(j)].width = 120 if (name == "Legend" and j == 2) else w
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 40
    wb.save(OUT)
    RESTRICTED.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, RESTRICTED / OUT.name)
    print(d[["Gene", "Literature evidence class (annotation)"] + gates + ["Core"]].to_string(index=False))


if __name__ == "__main__":
    main()
